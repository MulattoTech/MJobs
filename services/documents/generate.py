"""Generate factual, versioned local documents before application preparation."""
from dataclasses import asdict
from contextlib import contextmanager
import hashlib
import json
from pathlib import Path
import shutil
from uuid import uuid4

from docx import Document
from docx.shared import Pt, RGBColor
from pypdf import PdfReader
from sqlalchemy import select

from packages.domain.safety import Fact, select_verified_facts
from packages.persistence import models as m
from packages.persistence.store import identity, now, upsert
from services.ai.prompt_provenance import capture_document
from services.career_profile import story_selection
from services.dedupe.routes import persist_route_choices, plan_application_routes, qualify_route_authority
from services.dedupe.source_values import require_route_review
from services.applications.knowledge import application_adapter_knowledge
from services.documents import tailor
from services.documents.avoid_list import validate_content as validate_avoid_list
from services.documents.artifacts import atomic_generate, publish_directory
from services.documents.cover_letter import COVER_LETTER_TEMPLATE, build_cover_letter
from services.documents.errors import PreparationError
from services.documents.failure_status import DOCUMENT_GENERATION_FAILED, validate_generated_outputs
from services.documents.history import application_documents
from services.documents.page_fit import fit_to_page_target
from services.documents.render import render_pdf
from services.documents.validation import VALIDATOR, validate_assertions
from services.documents.templates import BASE_TEMPLATE, resume_layout, template_pair, contact_header, RESPONSIBILITY_EMPHASIS_TEMPLATES


def compose_paragraphs(profile, job, resume_rows, texts, template=BASE_TEMPLATE):
    return {'resume': resume_layout(profile, resume_rows, template)[0],
            'cover-letter': [profile['name'], f"Application for {job['title']}", *texts]}


@contextmanager
def uncommitted_files():
    paths = []
    try:
        yield paths
    except BaseException:
        for path in reversed(paths):
            if path.is_dir() and not path.is_symlink():
                shutil.rmtree(path, ignore_errors=True)
            else:
                path.unlink(missing_ok=True)
        raise


def _generate_document_artifact(destination, writer):
    """Translate renderer and writer failures while preserving recovery controls."""
    from workers.operations import OperationCancelled, OperationTimedOut

    try:
        return atomic_generate(destination, writer)
    except (PreparationError, OperationCancelled, OperationTimedOut):
        raise
    except Exception as error:
        raise PreparationError(
            DOCUMENT_GENERATION_FAILED,
            "Document generation failed validation. Generate a new revision before using this file.",
        ) from error


def _measure_and_fit_resume_rows(profile, job, rows, template, page_target, probe_dir,
                                 include_bullet_emphasis, explicitly_selected_ids=None,
                                 selected_stories=None):
    bullet_rows = [row for row in rows if row['kind'] in tailor.TARGET_BULLET_FACT_KINDS]
    skill_matches = tailor.matched_labels(bullet_rows, job)
    responsibility_matches = tailor.matched_responsibility_phrases(
        bullet_rows, tailor.target_responsibilities(job))
    explicitly_selected_ids = set(explicitly_selected_ids or ())
    fit_facts = []
    for row in rows:
        kind = row['kind']
        rank = (2 * len(skill_matches.get(row['fact_id'], ())) +
                len(responsibility_matches.get(row['fact_id'], ())))
        fit_facts.append({
            'fact_id': row['fact_id'],
            'kind': kind,
            'text': row['text'],
            'approved': bool(row['verified'] and row['user_approved']),
            'protected': kind in {
                'summary', 'experience', 'education', 'certification',
                'skill', 'technology_experience',
            },
            'optional_bullet': kind in {'responsibility', 'project', 'accomplishment'},
            'rank': rank,
            'explicitly_selected': row['fact_id'] in explicitly_selected_ids,
        })

    measurement_index = 0

    def measure_resume_pdf(candidate_facts):
        nonlocal measurement_index
        measurement_index += 1
        candidate_ids = {row['fact_id'] for row in candidate_facts}
        candidate_rows = [row for row in rows if row['fact_id'] in candidate_ids]
        candidate_paragraphs, candidate_headings = resume_layout(profile, candidate_rows, template)
        story_paragraphs, story_headings = story_selection.resume_story_paragraphs(selected_stories)
        story_offset = len(candidate_paragraphs)
        candidate_paragraphs.extend(story_paragraphs)
        candidate_headings.update(story_offset + index for index in story_headings)
        candidate_emphasis = tailor.emphasis(
            candidate_rows, job, include_responsibilities=include_bullet_emphasis)
        contact_count = len(contact_header(profile, template))
        candidate_fact_indices = [index for index in range(contact_count,
            len(candidate_paragraphs)) if index not in candidate_headings]
        candidate_bold_indices = {index for index, row in zip(candidate_fact_indices, candidate_rows)
            if row['fact_id'] in candidate_emphasis}
        probe_path = probe_dir / f'.resume-page-fit-attempt-{measurement_index}.pdf'
        _generate_document_artifact(probe_path, lambda path: render_pdf(
            path, candidate_paragraphs, heading_indices=candidate_headings,
            bold_indices=candidate_bold_indices))
        try:
            return len(PdfReader(str(probe_path)).pages)
        finally:
            probe_path.unlink(missing_ok=True)

    fit_result = fit_to_page_target(
        fit_facts, page_target=page_target, measure_page_count=measure_resume_pdf)
    selected_ids = {row['fact_id'] for row in fit_result.selected_facts}
    selected_rows = [row for row in rows if row['fact_id'] in selected_ids]
    return selected_rows, {
        'target_page_count': fit_result.page_target,
        'actual_page_count': fit_result.page_count,
        'unmet_target': fit_result.measurement_available and not fit_result.target_met,
        'measurement_status': 'MEASURED' if fit_result.measurement_available else 'UNAVAILABLE',
        'attempts_used': fit_result.attempts,
        'max_attempts': 8,
        'reason': fit_result.reason,
        'omitted_facts': list(fit_result.omitted_facts),
    }


def _page_fit_api_result(page_fit):
    return {
        'page_target': page_fit['target_page_count'],
        'page_count': page_fit['actual_page_count'],
        'attempts': page_fit['attempts_used'],
        'measurement_available': page_fit['measurement_status'] == 'MEASURED',
        'target_met': not page_fit['unmet_target'] and page_fit['measurement_status'] == 'MEASURED',
        'reason': page_fit['reason'],
        'omitted_facts': list(page_fit['omitted_facts']),
    }


def preview_resume_page_fit(store, profile_id, job_id, output_dir, target_pdf_pages,
                            resume_story_ids=None, resume_story_job_fingerprint=None,
                            expected_profile_version=None):
    """Render a transient measured PDF preview without creating an application or revision."""
    if type(target_pdf_pages) is not int or not 1 <= target_pdf_pages <= 20:
        raise ValueError('target_pdf_pages must be an integer from 1 through 20')
    profile = store.profile(profile_id)
    job = store.job(job_id, profile_id)
    selected_stories = []
    if resume_story_ids is not None:
        if (expected_profile_version != profile['version']
                or resume_story_job_fingerprint != tailor.target_fingerprint(job)):
            raise PreparationError('STORY_SELECTION_CONTEXT_STALE',
                'The profile or target job changed; reload before previewing selected stories')
        from services.career_profile.story_selection import StorySelectionError, select_stories_for_job
        try:
            selected_stories = select_stories_for_job(
                store, profile_id, job_id, 'RESUME', resume_story_ids,
            )['stories']
        except StorySelectionError as error:
            raise PreparationError(error.code, str(error)) from error
        if any(story['profile_version'] != profile['version'] for story in selected_stories):
            raise PreparationError('STORY_SELECTION_CONTEXT_STALE',
                'The profile changed; reload before previewing selected stories')
    eligible = tailor.eligible_rows(profile)
    if not eligible:
        raise PreparationError('DOCUMENT_NO_APPROVED_FACTS',
            'Approve and verify at least one sourced career fact before generating documents')
    rows = tailor.select_resume_rows(eligible, True, None)
    rows, _ = tailor.select_summary_rows(rows, job)
    rows, _ = tailor.select_bullet_rows(rows, job)
    template, _model = template_pair(True, False, False)
    include_bullet_emphasis = template in RESPONSIBILITY_EMPHASIS_TEMPLATES
    rows, _context = tailor.order_skills(rows, job,
        include_responsibilities=include_bullet_emphasis)

    preview_root = Path(output_dir).resolve() / '.page-fit-previews'
    if preview_root.is_symlink():
        raise PreparationError('DOCUMENT_PREVIEW_STORAGE_UNSAFE', 'Local preview storage is unavailable')
    preview_root.mkdir(parents=True, exist_ok=True)
    preview_id = uuid4().hex
    preview_pdf_path = preview_root / f'{preview_id}.pdf'
    selected_rows, page_fit = _measure_and_fit_resume_rows(
        profile, job, rows, template, target_pdf_pages, preview_root,
        include_bullet_emphasis, selected_stories=selected_stories)
    paragraphs, headings = resume_layout(profile, selected_rows, template)
    story_paragraphs, story_headings = story_selection.resume_story_paragraphs(selected_stories)
    story_offset = len(paragraphs)
    paragraphs.extend(story_paragraphs)
    headings.update(story_offset + index for index in story_headings)
    validate_avoid_list('\n'.join(paragraphs), profile.get('preferences'), 'resume')
    emphasized = tailor.emphasis(selected_rows, job,
        include_responsibilities=include_bullet_emphasis)
    fact_indices = [index for index in range(len(contact_header(profile, template)),
        len(paragraphs)) if index not in headings]
    bold_indices = {index for index, row in zip(fact_indices, selected_rows)
        if row['fact_id'] in emphasized}
    _generate_document_artifact(preview_pdf_path, lambda path: render_pdf(
        path, paragraphs, heading_indices=headings, bold_indices=bold_indices))
    return {
        'page_fit': _page_fit_api_result(page_fit),
        'preview_pdf_path': str(preview_pdf_path),
        'profile_version': profile['version'],
        'job_fingerprint': tailor.target_fingerprint(job),
        'resume_story_ids': list(resume_story_ids) if resume_story_ids is not None else None,
        'resume_story_job_fingerprint': resume_story_job_fingerprint,
    }


def discard_resume_page_fit_preview(output_dir, preview):
    """Remove only a PDF preview created inside this workspace's private preview directory."""
    raw_path = preview.get('preview_pdf_path') if isinstance(preview, dict) else None
    if not isinstance(raw_path, str):
        return
    preview_root = Path(output_dir).resolve() / '.page-fit-previews'
    path = Path(raw_path)
    if (preview_root.is_symlink() or path.parent != preview_root or len(path.stem) != 32 or
            any(char not in '0123456789abcdef' for char in path.stem) or path.suffix != '.pdf'):
        return
    path.unlink(missing_ok=True)
    try:
        preview_root.rmdir()
    except OSError:
        pass


def generate_resume_page_fit_preview(store, profile_id, job_id, output_dir, preview):
    """Persist a transient preview only after the API has collected explicit review."""
    if not isinstance(preview, dict) or not isinstance(preview.get('page_fit'), dict):
        raise ValueError('resume page-fit preview is invalid')
    page_fit = preview['page_fit']
    try:
        if not page_fit.get('measurement_available'):
            raise ValueError('resume page-fit preview has no measured PDF page count')
        if type(page_fit.get('page_target')) is not int or not 1 <= page_fit['page_target'] <= 20:
            raise ValueError('resume page-fit preview target is invalid')
        profile_version = preview.get('profile_version')
        job_fingerprint = preview.get('job_fingerprint')
        if not isinstance(profile_version, str) or not profile_version or not isinstance(job_fingerprint, str):
            raise ValueError('resume page-fit preview provenance is invalid')
        current_job = store.job(job_id, profile_id)
        if tailor.target_fingerprint(current_job) != job_fingerprint:
            raise ValueError('resume page-fit preview target changed')
        application_id = prepare(
            store, job_id, profile_id, output_dir, profile_version=profile_version,
            tailor_resume=True, resume_page_target=page_fit['page_target'],
            resume_page_fit_job_fingerprint=job_fingerprint,
            resume_story_ids=preview.get('resume_story_ids'),
            resume_story_job_fingerprint=preview.get('resume_story_job_fingerprint'))
        rows = application_documents(store, application_id)
        resume = next(row for row in rows if row['kind'] == 'resume' and
                      row['version'] == store.application(application_id)['document_version'])
        persisted = resume['generation_context'].get('page_fit')
        if _page_fit_api_result(persisted) != page_fit:
            raise ValueError('resume page-fit output no longer matches its preview')
        return {'application_id': application_id}
    finally:
        discard_resume_page_fit_preview(output_dir, preview)


def _resolved_role_lens(store, profile, job, profile_id, role_lens_id,
                        expected_role_lens_version, eligible):
    if role_lens_id is None and expected_role_lens_version is None:
        return None
    if role_lens_id is None or expected_role_lens_version is None:
        raise PreparationError(
            "ROLE_LENS_VERSION_INVALID",
            "Select a supported role lens and its current version together",
        )
    from services.career_profile.role_lenses import RoleLensError, select_role_lens
    try:
        selection = select_role_lens(
            store, profile_id, job.get("title"), role_lens_id=role_lens_id,
            expected_role_lens_version=expected_role_lens_version,
        )
    except RoleLensError as error:
        raise PreparationError(error.code, str(error)) from error

    eligible_ids = {row["fact_id"] for row in eligible}
    fact_ids = selection.get("fact_ids")
    emphasized_fact_ids = selection.get("emphasized_fact_ids")
    story_ids = selection.get("story_ids")
    emphasized_story_ids = selection.get("emphasized_story_ids")
    stories = selection.get("stories")
    source_refs = selection.get("source_refs")
    if (selection.get("profile_id") != profile_id
            or selection.get("profile_version") != profile.get("version")
            or not isinstance(fact_ids, list)
            or len(set(fact_ids)) != len(fact_ids)
            or not set(fact_ids).issubset(eligible_ids)
            or not isinstance(emphasized_fact_ids, list)
            or not set(emphasized_fact_ids).issubset(set(fact_ids))
            or not isinstance(story_ids, list)
            or len(set(story_ids)) != len(story_ids)
            or not isinstance(emphasized_story_ids, list)
            or not set(emphasized_story_ids).issubset(set(story_ids))
            or not isinstance(stories, list)
            or [row.get("story_id") for row in stories if isinstance(row, dict)] != story_ids
            or not isinstance(source_refs, dict)
            or [row.get("fact_id") for row in source_refs.get("facts", [])
                if isinstance(row, dict)] != fact_ids
            or [row.get("story_id") for row in source_refs.get("stories", [])
                if isinstance(row, dict)] != story_ids):
        raise PreparationError(
            "ROLE_LENS_CONTEXT_STALE",
            "Role-lens evidence changed; reload the profile and job before preparing documents",
        )
    lens = selection.get("lens")
    if (not isinstance(lens, dict) or lens.get("id") != role_lens_id
            or lens.get("version") != expected_role_lens_version
            or not isinstance(lens.get("display_name"), str)
            or not lens["display_name"].strip()):
        raise PreparationError(
            "ROLE_LENS_CONTEXT_STALE",
            "Role-lens evidence changed; reload the profile and job before preparing documents",
        )

    return {
        "schema_version": selection["schema_version"],
        "id": lens["id"],
        "version": lens["version"],
        "display_name": lens["display_name"],
        "description": lens.get("description"),
        "selection_reason": selection.get("selection_reason"),
        "profile_id": profile_id,
        "profile_version": profile["version"],
        "job_id": job["id"],
        "job_fingerprint": tailor.target_fingerprint(job),
        "fact_ids": list(fact_ids),
        "emphasized_fact_ids": list(emphasized_fact_ids),
        "story_ids": list(story_ids),
        "emphasized_story_ids": list(emphasized_story_ids),
        "source_refs": {
            "facts": [dict(row) for row in source_refs["facts"]],
            "stories": [dict(row) for row in source_refs["stories"]],
        },
    }


def prepare(store, job_id, profile_id, output_dir, check=None, profile_version=None, route_mode="DEFAULT", route_id=None, tailor_resume=False, tailor_fact_ids=None, tailor_fact_rewrites=None, include_resume_contact=False, cover_letter_style="professional", cover_letter_length="normal", resume_page_target=None, resume_page_fit_job_fingerprint=None, role_lens_id=None, expected_role_lens_version=None, cover_letter_story_ids=None, cover_letter_story_job_fingerprint=None, resume_story_ids=None, resume_story_job_fingerprint=None, expected_job_fingerprint=None, ai_evaluation_run_id=None):
    if type(tailor_resume) is not bool:
        raise ValueError('tailor_resume must be a boolean')
    if type(include_resume_contact) is not bool:
        raise ValueError('include_resume_contact must be a boolean')
    if resume_page_target is not None and (type(resume_page_target) is not int or
            not 1 <= resume_page_target <= 20):
        raise ValueError('resume_page_target must be an integer from 1 through 20')
    if resume_page_target is not None and not tailor_resume:
        raise ValueError('resume_page_target requires tailor_resume=true')
    if (resume_page_fit_job_fingerprint is not None and
            (not isinstance(resume_page_fit_job_fingerprint, str) or
             not resume_page_fit_job_fingerprint)):
        raise ValueError('resume_page_fit_job_fingerprint must be a nonempty string')
    resume_template,resume_model=template_pair(tailor_resume,tailor_fact_ids is not None,include_resume_contact,
        user_rephrase=tailor_fact_rewrites is not None)
    check = check or (lambda: None)
    with uncommitted_files() as staged, store.engine.begin() as connection:
        # A profile row serializes preparation and submission reservations.
        # SQLite has no row locks; use its explicit write transaction in tests.
        if connection.dialect.name == "sqlite":
            connection.exec_driver_sql("BEGIN IMMEDIATE")
        current = connection.execute(select(m.profiles).where(m.profiles.c.id == profile_id).with_for_update()).mappings().one()
        if profile_version and profile_version != current["version"]:
            raise PreparationError("DOCUMENT_PROFILE_CHANGED", "Profile changed; prepare documents using its current version")
        profile = store.profile(profile_id, profile_version)
        eligible = tailor.eligible_rows(profile)
        if not eligible:
            raise PreparationError("DOCUMENT_NO_APPROVED_FACTS", "Approve and verify at least one sourced career fact before generating documents")
        resume_rows = tailor.select_resume_rows(eligible, tailor_resume, tailor_fact_ids)
        selection_ids = [row['fact_id'] for row in resume_rows] if tailor_fact_ids is not None else None
        selection_hash = tailor.selection_fingerprint(selection_ids) if selection_ids is not None else None
        rewritten_rows, rephrase_context = (tailor.validate_user_rewrites(
            resume_rows, selection_ids, tailor_fact_rewrites) if tailor_fact_rewrites is not None
            else (resume_rows, []))
        rephrase_hash = hashlib.sha256(json.dumps(tailor_fact_rewrites,
            ensure_ascii=False, sort_keys=True, separators=(',',':')).encode()).hexdigest() if tailor_fact_rewrites is not None else None
        check()
        job = dict(connection.execute(select(m.jobs, m.companies.c.name.label("company")).join(
            m.companies, m.companies.c.id == m.jobs.c.company_id).where(m.jobs.c.id == job_id)
            .with_for_update(of=m.jobs)).mappings().one())
        selected_resume_stories = []
        if ai_evaluation_run_id is not None:
            try:
                store.require_reviewed_ai_document_candidate(
                    profile_id, ai_evaluation_run_id, job_id, connection=connection)
            except ValueError as error:
                raise PreparationError(str(error), 'AI evaluation inputs changed; evaluate and review again') from None
        if expected_job_fingerprint is not None and tailor.target_fingerprint(job) != expected_job_fingerprint:
            raise PreparationError("DOCUMENT_JOB_CHANGED", "Job changed; review it before preparing again")
        if resume_story_ids is not None:
            if (not isinstance(resume_story_ids, list) or not resume_story_ids
                    or len(set(resume_story_ids)) != len(resume_story_ids)):
                raise PreparationError(
                    "STORY_SELECTION_INVALID",
                    "Choose one or more distinct reviewed stories or omit resume story selection",
                )
            if (not isinstance(resume_story_job_fingerprint, str)
                    or not resume_story_job_fingerprint
                    or tailor.target_fingerprint(job) != resume_story_job_fingerprint):
                raise PreparationError(
                    "STORY_SELECTION_CONTEXT_STALE",
                    "The target job changed; reload it and review story selection again",
                )
            from services.career_profile.story_selection import StorySelectionError, select_stories_for_job
            try:
                selected_resume_stories = select_stories_for_job(
                    store, profile_id, job_id, "RESUME", resume_story_ids,
                    connection=connection,
                )["stories"]
            except StorySelectionError as error:
                raise PreparationError(error.code, str(error)) from error
            if any(story["profile_version"] != profile["version"]
                   for story in selected_resume_stories):
                raise PreparationError(
                    "STORY_SELECTION_CONTEXT_STALE",
                    "Career facts changed; reload the profile and review story selection again",
                )
        resume_story_refs = story_selection.story_revision_refs(selected_resume_stories)
        selected_stories = []
        if cover_letter_story_ids is not None:
            if (not isinstance(cover_letter_story_ids, list)
                    or not cover_letter_story_ids
                    or len(set(cover_letter_story_ids)) != len(cover_letter_story_ids)):
                raise PreparationError(
                    "STORY_SELECTION_INVALID",
                    "Choose one or more distinct reviewed stories or omit story selection",
                )
            if (not isinstance(cover_letter_story_job_fingerprint, str)
                    or not cover_letter_story_job_fingerprint
                    or tailor.target_fingerprint(job) != cover_letter_story_job_fingerprint):
                raise PreparationError(
                    "STORY_SELECTION_CONTEXT_STALE",
                    "The target job changed; reload it and review story selection again",
                )
            from services.career_profile.story_selection import (
                StorySelectionError,
                select_stories_for_job,
            )
            try:
                selection = select_stories_for_job(
                    store, profile_id, job_id, "COVER_LETTER",
                    cover_letter_story_ids, connection=connection,
                )
            except StorySelectionError as error:
                raise PreparationError(error.code, str(error)) from error
            selected_stories = selection["stories"]
            if any(story["profile_version"] != profile["version"]
                   for story in selected_stories):
                raise PreparationError(
                    "STORY_SELECTION_CONTEXT_STALE",
                    "Career facts changed; reload the profile and review story selection again",
                )
        if (resume_page_fit_job_fingerprint is not None and
                tailor.target_fingerprint(job) != resume_page_fit_job_fingerprint):
            raise PreparationError('RESUME_PAGE_FIT_PREVIEW_STALE_JOB',
                'The target job changed; preview page fit again before generating documents')
        from services.company_discovery.company_context import approved_company_context_for_job
        job_attributes = dict(job.get("attributes") or {})
        job_attributes["approved_company_context"] = approved_company_context_for_job(
            connection, job_id, job["company_id"])
        job["attributes"] = job_attributes
        role_lens = _resolved_role_lens(
            store, profile, job, profile_id, role_lens_id,
            expected_role_lens_version, eligible,
        )
        try:
            require_route_review(connection, job_id, route_mode, route_id)
        except ValueError as error:
            raise PreparationError(str(error), 'The selected source link differs from a preserved application route; review and select the route explicitly') from error
        routes = qualify_route_authority(connection, [dict(row) for row in connection.execute(
            select(m.routes).where(m.routes.c.job_id == job_id)
            .order_by(m.routes.c.id)).mappings()])
        route_preference = (profile.get("preferences") or {}).get(
            "application_route_preference", "PREFER_VERIFIED")
        try:
            route_plan = plan_application_routes(
                routes, route_mode, route_id, route_preference)
        except ValueError as error:
            raise PreparationError(str(error), "Choose an application route before preparing documents") from error
        route = route_plan["selected"][0]["preferred_route"]
        cover_letter_paragraphs, cover_letter_context = build_cover_letter(
            profile, job, profile["facts"], cover_letter_style, cover_letter_length,
            selected_stories=selected_stories)
        if role_lens:
            cover_letter_context["role_lens"] = role_lens
            cover_letter_context["job_snapshot"]["role_lens"] = role_lens
        prior = connection.execute(select(m.applications).where(m.applications.c.job_id == job_id, m.applications.c.profile_id == profile_id).with_for_update()).mappings().one_or_none()
        if prior:
            prior_resume=connection.execute(select(m.documents).where(m.documents.c.application_id==prior['id'],
                m.documents.c.version==prior['document_version'],m.documents.c.kind=='resume')).mappings().one_or_none()
            expected_template=resume_template
            same_mode=prior_resume and prior_resume['template']==expected_template
            same_target=not tailor_resume or (prior_resume and (prior_resume['generation_context'] or {}).get('target_fingerprint')==tailor.target_fingerprint(job))
            same_selection=selection_ids is None or (prior_resume and (prior_resume['generation_context'] or {}).get('selection_fingerprint')==selection_hash)
            same_rephrases=(rephrase_hash is None or (prior_resume and
                (prior_resume['generation_context'] or {}).get('rephrase_fingerprint')==rephrase_hash))
            prior_page_fit=(prior_resume['generation_context'] or {}).get('page_fit') if prior_resume else None
            same_page_fit=((resume_page_target is None and prior_page_fit is None) or
                (resume_page_target is not None and prior_page_fit is not None and
                 prior_page_fit.get('target_page_count') == resume_page_target))
            prior_resume_context = (prior_resume['generation_context'] or {}) if prior_resume else {}
            same_resume_stories = bool(prior_resume and
                prior_resume_context.get('story_refs', []) == resume_story_refs and
                prior_resume_context.get('story_job_id') == (job_id if resume_story_refs else None) and
                prior_resume_context.get('story_job_fingerprint') ==
                    (resume_story_job_fingerprint if resume_story_refs else None))
            prior_letter=connection.execute(select(m.documents).where(m.documents.c.application_id==prior['id'],
                m.documents.c.version==prior['document_version'],m.documents.c.kind=='cover-letter')).mappings().one_or_none()
            prior_letter_context = (prior_letter['generation_context'] or {}) if prior_letter else {}
            same_letter=bool(prior_letter and prior_letter['template']==COVER_LETTER_TEMPLATE
                and {key:value for key,value in prior_letter_context.items()
                     if key not in {'claim_validation', 'prompt_provenance'}} == cover_letter_context)
            if prior["profile_version"] == profile["version"] and same_mode and same_target and same_selection and same_rephrases and same_page_fit and same_resume_stories and same_letter and prior['state'] != 'REVIEW_REQUIRED':
                persist_route_choices(connection, prior["id"], route_plan)
                return prior["id"]
            if prior["state"] not in ("PREPARED", "DRY_RUN_COMPLETE", "REVIEW_REQUIRED"):
                raise PreparationError("DOCUMENT_HISTORY_LOCKED", "Historical or active application cannot be regenerated")
        selected = select_verified_facts([Fact(id=row["fact_id"], **{key: row[key] for key in ("text", "source", "verified", "user_approved", "confidence")})
                                          for row in eligible], [row["fact_id"] for row in eligible])
        application_id = identity(job_id, profile_id)
        revision = prior["document_version"] + 1 if prior else 1
        relative = profile["version"] + "/" + str(revision)
        root = Path(output_dir).resolve() / application_id / relative
        staging = root.with_name(f".{root.name}.staging-{uuid4().hex}")
        staging.mkdir(parents=True, exist_ok=False)
        staged.append(staging)
        fact_ids = [fact.id for fact in selected]
        texts = [fact.text for fact in selected]
        summary_adjustment=None
        bullet_adjustment=None
        if tailor_resume and tailor_fact_ids is None:
            resume_rows,summary_adjustment=tailor.select_summary_rows(resume_rows,job)
            resume_rows,bullet_adjustment=tailor.select_bullet_rows(resume_rows,job)
        include_bullet_emphasis=resume_template in RESPONSIBILITY_EMPHASIS_TEMPLATES
        ordered_rows,context=tailor.order_skills(resume_rows,job,
            include_responsibilities=include_bullet_emphasis) if tailor_resume else (resume_rows,None)
        emphasized=tailor.emphasis(ordered_rows,job,
            include_responsibilities=include_bullet_emphasis) if tailor_resume else []
        if context is not None:
            context['emphasized_fact_ids']=emphasized
            if summary_adjustment is not None:
                context['summary_adjustment']=summary_adjustment
            if bullet_adjustment is not None:
                context['bullet_adjustment']=bullet_adjustment
        if selection_ids is not None:
            context.update(selection_mode='EXPLICIT_APPROVED_FACTS',selected_fact_ids=selection_ids,
                available_fact_ids=fact_ids,omitted_fact_ids=[value for value in fact_ids if value not in set(selection_ids)],
                selection_fingerprint=selection_hash)
        if rephrase_context:
            ordered_by_id={row['fact_id']:row for row in rewritten_rows}
            ordered_rows=[dict(row, text=ordered_by_id[row['fact_id']]['text']) for row in ordered_rows]
            context['fact_rephrases']=rephrase_context
            context['rephrase_fingerprint']=rephrase_hash
        if resume_page_target is not None:
            ordered_rows, page_fit = _measure_and_fit_resume_rows(
                profile, job, ordered_rows, resume_template, resume_page_target,
                staging, include_bullet_emphasis, selection_ids,
                selected_stories=selected_resume_stories)
            selected_fit_ids = {row['fact_id'] for row in ordered_rows}
            if context and 'emphasized_fact_ids' in context:
                context['emphasized_fact_ids'] = [
                    fact_id for fact_id in context['emphasized_fact_ids'] if fact_id in selected_fit_ids]
            context = context or {'mode': 'APPROVED_FACT_PAGE_FIT'}
            context['page_fit'] = page_fit
        if role_lens:
            rows_by_id = {row['fact_id']: row for row in ordered_rows}
            lens_order = [fact_id for fact_id in role_lens['fact_ids'] if fact_id in rows_by_id]
            ordered_rows = ([rows_by_id[fact_id] for fact_id in lens_order]
                            + [row for row in ordered_rows if row['fact_id'] not in set(lens_order)])
            included_ids = {row['fact_id'] for row in ordered_rows}
            emphasized = list(dict.fromkeys([
                *(fact_id for fact_id in emphasized if fact_id in included_ids),
                *(fact_id for fact_id in role_lens['emphasized_fact_ids']
                  if fact_id in included_ids),
            ]))
            role_lens['document_fact_ids'] = [row['fact_id'] for row in ordered_rows]
            role_lens['document_emphasized_fact_ids'] = list(emphasized)
            context = context or {'mode': 'ROLE_LENS_PRESENTATION'}
            context['ordered_fact_ids'] = [row['fact_id'] for row in ordered_rows]
            context['emphasized_fact_ids'] = emphasized
            context['role_lens'] = role_lens
        resume_ids=[row['fact_id'] for row in ordered_rows]
        if resume_story_refs:
            context = context or {'mode': 'APPROVED_FACTS'}
            context['story_refs'] = resume_story_refs
            context['review_required'] = True
            context['story_job_id'] = job_id
            context['story_job_fingerprint'] = resume_story_job_fingerprint
        if include_resume_contact:
            context=context or {'mode':'APPROVED_CONTACT_HEADER'}
            context['contact_header']=contact_header(profile,resume_template)
        resume_provenance = capture_document('resume', resume_template, resume_model)
        letter_provenance = capture_document('cover_letter', COVER_LETTER_TEMPLATE,
                                             'deterministic-job-evidence-v1')
        kinds = compose_paragraphs(profile, job, ordered_rows, texts, resume_template)
        resume_story_paragraphs, resume_story_headings = story_selection.resume_story_paragraphs(
            selected_resume_stories)
        resume_story_offset = len(kinds['resume'])
        kinds['resume'].extend(resume_story_paragraphs)
        kinds['cover-letter'] = cover_letter_paragraphs
        if set(kinds) != {'resume', 'cover-letter'}:
            raise PreparationError('UNSUPPORTED_CLAIM', 'The generated document set has no complete validation contract')
        for kind, paragraphs in kinds.items():
            output_path = 'resume' if kind == 'resume' else 'cover letter'
            validate_avoid_list('\n'.join(paragraphs), profile.get('preferences'), output_path)
        records = []
        generated_at = now()
        for kind, paragraphs in kinds.items():
            check()
            template = resume_template if kind=='resume' else COVER_LETTER_TEMPLATE
            current_fact_ids = resume_ids if kind=='resume' else cover_letter_context['fact_ids']
            generation_context = context if kind=='resume' else cover_letter_context
            claim_validation = validate_assertions(
                kind, paragraphs, profile, current_fact_ids, job['title'], template,
                generation_context=generation_context,
                selected_stories=(selected_stories if kind == 'cover-letter'
                                  else selected_resume_stories if kind == 'resume' else None))
            generation_context = {**(generation_context or {}),
                                 'claim_validation': claim_validation,
                                 'prompt_provenance': (resume_provenance if kind == 'resume'
                                                       else letter_provenance)}
            headings = resume_layout(profile, ordered_rows, template)[1] if kind=='resume' else set()
            if kind == 'resume':
                headings.update(resume_story_offset + index for index in resume_story_headings)
            fact_indices=[index for index in range(len(contact_header(profile,template)),len(paragraphs)) if index not in headings] if kind=='resume' else []
            bold_indices={index for index,row in zip(fact_indices,ordered_rows) if row['fact_id'] in emphasized}
            document = Document()
            document.styles['Normal'].font.name = 'Arial'
            document.styles['Normal'].font.size = Pt(11)
            document.styles['Heading 1'].font.name = 'Arial'
            document.styles['Heading 1'].font.size = Pt(13)
            document.styles['Heading 1'].font.color.rgb = RGBColor(0,0,0)
            for index,text in enumerate(paragraphs):
                paragraph=document.add_paragraph(style='Heading 1' if index in headings else 'Normal')
                run=paragraph.add_run(text)
                if index in bold_indices:run.bold=True
            docx_path = staging / f"{kind}.docx"
            pdf_path = staging / f"{kind}.pdf"
            txt_path = staging / f"{kind}.txt"
            _generate_document_artifact(docx_path, document.save)
            _generate_document_artifact(txt_path, lambda path: path.write_text(
                '\n'.join(paragraphs) + '\n', encoding='utf-8', newline='\n'))
            _generate_document_artifact(pdf_path, lambda path: render_pdf(
                path, paragraphs, heading_indices=headings, bold_indices=bold_indices))
            validate_generated_outputs({'docx': docx_path, 'pdf': pdf_path, 'txt': txt_path})
            records.append({"id": identity(application_id, kind, revision), "application_id": application_id,
                "kind": kind, "version": revision, "profile_version": profile["version"],
                "files": {path.suffix[1:]: relative + '/' + path.name for path in (docx_path, pdf_path, txt_path)},
                "fact_ids": current_fact_ids, "sha256": {path.suffix[1:]: hashlib.sha256(path.read_bytes()).hexdigest() for path in (docx_path, pdf_path, txt_path)},
                "template": template, "created_at": generated_at,
                "generation_model": resume_model if kind=='resume' else "deterministic-job-evidence-v1",
                "generation_context": generation_context})
        check()
        publish_directory(staging, root)
        staged.append(root)
        upsert(connection, m.applications, {"id": application_id, "job_id": job_id, "route_id": route["id"], "profile_id": profile_id,
            "profile_version": profile["version"], "document_version": revision, "state": "PREPARED", "target": "MOCK" if route["provider"] == "mock-ats" else "REAL", "fact_ids": fact_ids,
            "trace": [*(prior["trace"] if prior else []), "GENERATE_DOCUMENTS", "PREPARE_APPLICATION"], "confirmation": None, "created_at": prior["created_at"] if prior else now()}, replace=bool(prior))
        selected_route_ids = persist_route_choices(connection, application_id, route_plan)
        for record in records:
            upsert(connection, m.documents, record)
        store._audit(connection, application_id, 'DOCUMENTS_VALIDATED', {'validator': VALIDATOR,
            'profile_version': profile['version'], 'documents': [
                {'id': r['id'], 'fact_ids': r['fact_ids'], 'sha256': r['sha256'],
                 'claim_validation': r['generation_context']['claim_validation']}
                for r in records]})
        store._audit(connection, application_id, "DOCUMENTS_PREPARED", {"fact_ids": fact_ids, "profile_version": profile["version"], "route_ids": selected_route_ids, "route_mode": route_plan["mode"],
            "document_kinds": sorted({record["kind"] for record in records}),
            "adapter_knowledge": application_adapter_knowledge(connection, route["id"])})
        return application_id
