"""Generation and automatic use reject unsupported or corrupted assertions."""
import hashlib
import unittest
from unittest.mock import patch
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from sqlalchemy import select, update

from fixtures.mock_server import fixture_server
from packages.persistence import models as m
from services.discovery.pipeline import discover
from services.documents import generate
from services.documents.history import application_documents
from services.applications.mock_adapter import apply_mock
from tests.integration import test_document_provenance as fixtures
from tests.integration.test_spine import PROFILE


class DocumentValidationTests(unittest.TestCase):
    setUp = fixtures.DocumentProvenanceTests.setUp
    tearDown = fixtures.DocumentProvenanceTests.tearDown

    def test_unicode_readback_and_validation_receipts(self):
        self.store.save_profile(dict(PROFILE, facts=[dict(PROFILE['facts'][0], text='Δοκιμή Python Система résumé')]))
        with fixture_server() as origin:
            discover(self.store, origin, 'fixture')
            app = generate.prepare(self.store, self.store.jobs()[0]['id'], 'fixture', self.root)
        for row in application_documents(self.store, app):
            paths = {ext: self.root/app/file for ext,file in row['files'].items()}
            expected = paths['txt'].read_text(encoding='utf-8').split()
            self.assertEqual(' '.join(p.text for p in Document(paths['docx']).paragraphs).split(), expected)
            self.assertEqual(' '.join(p.extract_text() for p in PdfReader(paths['pdf']).pages).split(), expected)
        with self.store.engine.connect() as c:
            receipts = c.execute(select(m.audit).where(m.audit.c.entity_id==app,m.audit.c.action=='DOCUMENTS_VALIDATED')).mappings().all()
        self.assertEqual(len(receipts), 1)
        self.assertEqual(receipts[0]['payload']['validator'], 'exact-approved-facts-v5-claim-preservation')
        self.assertEqual(len(receipts[0]['payload']['documents']), 2)
        for row in application_documents(self.store, app):
            claim_validation = row['generation_context']['claim_validation']
            self.assertEqual(claim_validation['strategy'], 'DETERMINISTIC_EXACT_FACT_PRESERVATION')
            self.assertEqual(claim_validation['keep_drop_authority'], 'TEMPLATE_AND_USER_SELECTION')
            self.assertTrue(claim_validation['claims'])
        for record in receipts[0]['payload']['documents']:
            self.assertIn('claim_validation', record)

    def test_claim_ledger_is_profile_scoped_persisted_and_redacted(self):
        facts = [
            {'id': 'work', 'kind': 'experience',
             'text': 'Applied Python at Northstar Systems from 2020 to 2024; reduced processing time 35 percent.',
             'source': 'synthetic-test-profile', 'verified': True,
             'user_approved': True, 'confidence': 100},
            {'id': 'degree', 'kind': 'education',
             'text': 'Bachelor of Science, North Valley University, 2018.',
             'source': 'synthetic-test-profile', 'verified': True,
             'user_approved': True, 'confidence': 100},
        ]
        self.store.save_profile(dict(PROFILE, facts=facts))
        with fixture_server() as origin:
            discover(self.store, origin, 'fixture')
            job = self.store.jobs()[0]
            application_id = generate.prepare(self.store, job['id'], 'fixture', self.root)
        resume = next(row for row in application_documents(self.store, application_id)
                      if row['kind'] == 'resume')
        evidence = resume['generation_context']['claim_validation']
        by_id = {claim['fact_id']: claim for claim in evidence['claims']}
        self.assertEqual(by_id['work']['kind'], 'experience')
        self.assertIn('EMPLOYER_TEXT_WITHIN_EXPERIENCE_FACT',
                      by_id['work']['protected_fields'])
        self.assertIn('DEGREE_TEXT_WITHIN_EDUCATION_FACT',
                      by_id['degree']['protected_fields'])
        self.assertEqual(by_id['work']['protected_token_counts'], {'dates': 2, 'numbers': 3})
        serialized_evidence = str(evidence)
        self.assertNotIn('Northstar Systems', serialized_evidence)
        self.assertNotIn('North Valley University', serialized_evidence)
        self.assertNotIn('2020', serialized_evidence)
        with self.store.engine.connect() as connection:
            receipt = connection.execute(select(m.audit).where(
                m.audit.c.entity_id == application_id,
                m.audit.c.action == 'DOCUMENTS_VALIDATED')).mappings().one()
        audit_claims = next(row['claim_validation'] for row in receipt['payload']['documents']
                            if row['id'] == resume['id'])
        self.assertEqual(audit_claims, evidence)

        other_facts = [dict(row) for row in facts]
        other_facts[0]['text'] = other_facts[0]['text'].replace(
            'Northstar Systems', 'Southstar Systems')
        self.store.save_profile(dict(PROFILE, id='other', name='Other Fixture', facts=other_facts))
        other_application = generate.prepare(self.store, job['id'], 'other', self.root)
        other_resume = next(row for row in application_documents(self.store, other_application)
                            if row['kind'] == 'resume')
        other_claims = {claim['fact_id']: claim for claim in
                        other_resume['generation_context']['claim_validation']['claims']}
        self.assertEqual(set(other_claims), set(by_id))
        self.assertEqual(other_claims['work']['decision'], 'PRESERVED_EXACT')
        self.assertNotEqual(other_resume['application_id'], resume['application_id'])
        first_text = (self.root / resume['application_id'] /
                      resume['files']['txt']).read_text(encoding='utf-8')
        other_text = (self.root / other_resume['application_id'] /
                      other_resume['files']['txt']).read_text(encoding='utf-8')
        self.assertIn('Northstar Systems', first_text)
        self.assertIn('Southstar Systems', other_text)

    def test_unsupported_claim_or_glyph_preserves_prior_documents(self):
        with fixture_server() as origin:
            discover(self.store, origin, 'fixture'); job=self.store.jobs()[0]
            app=generate.prepare(self.store, job['id'], 'fixture', self.root)
            before=application_documents(self.store,app); old_app=self.store.application(app)
            files={str(p):p.read_bytes() for p in (self.root/app).rglob('*') if p.is_file()}
            self.store.save_profile(dict(PROFILE, facts=[dict(PROFILE['facts'][0],text='Python systems')]))
            with patch('services.documents.generate.compose_paragraphs',return_value={'resume':['Alex Fixture','alex@example.test','Invented certification'], 'cover-letter':['Alex Fixture']}):
                with self.assertRaises(generate.PreparationError) as error:
                    generate.prepare(self.store,job['id'],'fixture',self.root)
            self.assertEqual(error.exception.code,'UNSUPPORTED_CLAIM')
            with patch('services.documents.generate.compose_paragraphs', return_value={
                    'resume': ['Alex Fixture', 'alex@example.test',
                               'Python at Northstar Systems from 2025 with a 95 percent success rate'],
                    'cover-letter': ['unused']}):
                with self.assertRaises(generate.PreparationError) as error:
                    generate.prepare(self.store, job['id'], 'fixture', self.root)
            self.assertEqual(error.exception.code, 'UNSUPPORTED_CLAIM')
            self.assertIn('date claim', str(error.exception))
            self.store.save_profile(dict(PROFILE, facts=[dict(PROFILE['facts'][0],text='Python 中文 systems')]))
            with self.assertRaises(generate.PreparationError) as error:
                generate.prepare(self.store,job['id'],'fixture',self.root)
            self.assertEqual(error.exception.code,'DOCUMENT_UNSUPPORTED_GLYPH')
            self.assertEqual(application_documents(self.store,app),before)
            self.assertEqual(self.store.application(app),old_app)
            self.assertEqual(files,{str(p):p.read_bytes() for p in (self.root/app).rglob('*') if p.is_file()})

    def test_legacy_hash_valid_unsupported_resume_never_reaches_browser(self):
        with fixture_server() as origin:
            discover(self.store,origin,'fixture')
            app=generate.prepare(self.store,self.store.jobs()[0]['id'],'fixture',self.root)
            row=next(r for r in application_documents(self.store,app) if r['kind']=='resume')
            path=self.root/app/row['files']['docx']
            doc=Document(path);doc.add_paragraph('Invented certification');doc.save(path)
            with self.store.engine.begin() as c:
                c.execute(update(m.documents).where(m.documents.c.id==row['id']).values(sha256={**row['sha256'],'docx':hashlib.sha256(path.read_bytes()).hexdigest()}))
            with patch('services.applications.mock_adapter.sync_playwright') as browser:
                with self.assertRaises(generate.PreparationError) as error:
                    apply_mock(self.store,app,origin,self.root)
                browser.assert_not_called()
            self.assertEqual(error.exception.code,'UNSUPPORTED_CLAIM')
            self.assertEqual(self.store.application(app)['state'],'REVIEW_REQUIRED')
            generate.prepare(self.store,self.store.application(app)['job_id'],'fixture',self.root)
            self.assertEqual(self.store.application(app)['state'],'PREPARED')
            self.assertEqual(self.store.application(app)['document_version'],2)
