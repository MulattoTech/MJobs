"""ATS structure preserves factual text, historical contracts and safe use."""
import hashlib
import unittest
from unittest.mock import patch
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
from sqlalchemy import update

from fixtures.mock_server import fixture_server
from packages.persistence import models as m
from services.discovery.pipeline import discover
from services.documents.generate import prepare
from services.documents.history import application_documents
from services.documents.comparison import compare_resumes
from services.documents.library import document_library
from services.documents.render import render_pdf
from services.documents.validation import validate_stored_resume
from services.applications.mock_adapter import apply_mock
from tests.integration import test_document_provenance as fixtures
from tests.integration.test_spine import PROFILE


class AtsTemplateTests(unittest.TestCase):
    setUp=fixtures.DocumentProvenanceTests.setUp
    tearDown=fixtures.DocumentProvenanceTests.tearDown

    def test_headings_are_structural_and_all_formats_preserve_fact_order(self):
        facts=[dict(PROFILE['facts'][0],id='summary',kind='summary',text='Systems practitioner.'),
            PROFILE['facts'][0],dict(PROFILE['facts'][0],id='role',kind='experience',text='Example Engineer 2020-2024'),
            dict(PROFILE['facts'][0],id='duty',kind='responsibility',text='Maintained services.'),
            PROFILE['facts'][1],dict(PROFILE['facts'][0],id='literal',kind='custom',text='Skills')]
        self.store.save_profile(dict(PROFILE,facts=facts))
        with fixture_server() as origin:discover(self.store,origin,'fixture')
        app=prepare(self.store,self.store.jobs()[0]['id'],'fixture',self.root)
        row=next(r for r in application_documents(self.store,app) if r['kind']=='resume')
        self.assertEqual(row['template'],'verified-headings-v1')
        paths={ext:self.root/app/file for ext,file in row['files'].items()}
        expected=['Alex Fixture','alex@example.test','Summary','Systems practitioner.','Skills','Python',
            'Experience','Example Engineer 2020-2024','Maintained services.','Skills','SQL','Additional Information','Skills']
        self.assertEqual(paths['txt'].read_text(encoding='utf-8').splitlines(),expected)
        doc=Document(paths['docx'])
        self.assertEqual([p.text for p in doc.paragraphs],expected)
        self.assertEqual([p.text for p in doc.paragraphs if p.style.name=='Heading 1'],['Summary','Skills','Experience','Skills','Additional Information'])
        self.assertEqual(doc.paragraphs[-1].style.name,'Normal')
        self.assertFalse(doc.tables);self.assertFalse(doc.inline_shapes)
        self.assertTrue(all(node.get(qn('w:num'),'1')=='1' for node in doc.element.xpath('.//w:cols')))
        self.assertNotIn(b'<w:drawing',doc.part.blob)
        self.assertEqual(' '.join(page.extract_text() for page in PdfReader(paths['pdf']).pages).split(),' '.join(expected).split())
        self.assertTrue(all(not page.images for page in PdfReader(paths['pdf']).pages))
        validate_stored_resume(paths,self.store.profile('fixture'),row['fact_ids'],row['template'],row['generation_model'])

    def test_legacy_revisions_remain_comparable_and_explicit_prepare_upgrades_once(self):
        with fixture_server() as origin:discover(self.store,origin,'fixture')
        job=self.store.jobs()[0];app=prepare(self.store,job['id'],'fixture',self.root)
        row=next(r for r in application_documents(self.store,app) if r['kind']=='resume')
        # Reconstruct a genuine legacy contract and its corresponding bytes.
        paragraphs=['Alex Fixture','alex@example.test',*[f['text'] for f in row['facts_used']]]
        paths={ext:self.root/app/file for ext,file in row['files'].items()}
        doc=Document()
        for value in paragraphs:doc.add_paragraph(value)
        doc.save(paths['docx']);paths['txt'].write_text('\n'.join(paragraphs)+'\n',encoding='utf-8',newline='\n');render_pdf(paths['pdf'],paragraphs)
        with self.store.engine.begin() as c:c.execute(update(m.documents).where(m.documents.c.id==row['id']).values(
            template='verified-facts-v1',generation_model=None,sha256={k:hashlib.sha256(p.read_bytes()).hexdigest() for k,p in paths.items()}))
        originals={str(p):p.read_bytes() for p in paths.values()}
        validate_stored_resume(paths,self.store.profile('fixture'),row['fact_ids'],'verified-facts-v1',None)
        prepare(self.store,job['id'],'fixture',self.root)
        self.assertEqual(self.store.application(app)['document_version'],2)
        prepare(self.store,job['id'],'fixture',self.root)
        self.assertEqual(self.store.application(app)['document_version'],2)
        prepare(self.store,job['id'],'fixture',self.root,tailor_resume=True)
        target=next(r for r in application_documents(self.store,app) if r['kind']=='resume' and r['version']==3)
        result=compare_resumes(self.store,self.root,app,row['id'],target['id'])
        self.assertEqual(result['omitted'],[])
        self.assertIn('Skills',result['tailored']['text'])
        self.assertEqual(next(r for r in document_library(self.store,'fixture')['generated'] if r['id']==target['id'])['document_kind'],'TAILORED_RESUME')
        for path,content in originals.items():self.assertEqual(__import__('pathlib').Path(path).read_bytes(),content)

    def test_heading_tampering_fails_before_browser(self):
        with fixture_server() as origin:
            discover(self.store,origin,'fixture');app=prepare(self.store,self.store.jobs()[0]['id'],'fixture',self.root)
            row=next(r for r in application_documents(self.store,app) if r['kind']=='resume')
            path=self.root/app/row['files']['docx'];doc=Document(path)
            self.assertEqual(doc.paragraphs[2].text,'Skills')
            doc.paragraphs[2].text='Certified Expert';doc.save(path)
            with self.store.engine.begin() as c:c.execute(update(m.documents).where(m.documents.c.id==row['id']).values(sha256={**row['sha256'],'docx':hashlib.sha256(path.read_bytes()).hexdigest()}))
            with patch('services.applications.mock_adapter.sync_playwright') as browser:
                with self.assertRaisesRegex(ValueError,'assertions'):
                    apply_mock(self.store,app,origin,self.root)
                browser.assert_not_called()
            self.assertEqual(self.store.application(app)['state'],'REVIEW_REQUIRED')
