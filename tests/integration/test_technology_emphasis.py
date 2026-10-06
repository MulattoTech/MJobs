"""Real serialized emphasis changes presentation, never approved assertions."""
import unittest
from docx import Document
from pypdf import PdfReader
from fixtures.mock_server import fixture_server
from services.discovery.pipeline import discover
from services.documents.generate import prepare
from services.documents.history import application_documents
from services.documents.comparison import compare_resumes
from tests.integration import test_document_provenance as fixtures
from tests.integration.test_spine import PROFILE


class TechnologyEmphasisTests(unittest.TestCase):
    setUp=fixtures.DocumentProvenanceTests.setUp
    tearDown=fixtures.DocumentProvenanceTests.tearDown

    def test_matched_approved_technologies_are_bold_in_both_formats_only(self):
        tech=dict(PROFILE['facts'][0],id='tech',kind='technology_experience',text='Built Python services & <tools>.')
        duty=dict(PROFILE['facts'][0],id='python-duty',kind='responsibility',text='Built Python APIs for internal teams.')
        unrelated=dict(PROFILE['facts'][0],id='unrelated-duty',kind='responsibility',text='Scheduled COBOL batch jobs.')
        self.store.save_profile(dict(PROFILE,facts=[*PROFILE['facts'],tech,duty,unrelated]))
        with fixture_server() as origin:discover(self.store,origin,'fixture')
        job=next(j for j in self.store.jobs() if 'Python' in j['skills'])
        app=prepare(self.store,job['id'],'fixture',self.root)
        base=next(r for r in application_documents(self.store,app) if r['kind']=='resume')
        original={ext:(self.root/app/path).read_bytes() for ext,path in base['files'].items()}
        before=self.store.profile('fixture')
        prepare(self.store,job['id'],'fixture',self.root,tailor_resume=True)
        row=next(r for r in application_documents(self.store,app) if r['kind']=='resume' and r['version']==2)
        self.assertIn('tech',row['generation_context']['emphasized_fact_ids'])
        self.assertIn('python-duty',row['generation_context']['emphasized_fact_ids'])
        self.assertNotIn('unrelated-duty',row['generation_context']['emphasized_fact_ids'])
        doc=Document(self.root/app/row['files']['docx'])
        bold=[p.text for p in doc.paragraphs if p.runs and all(r.bold for r in p.runs)]
        self.assertIn(tech['text'],bold);self.assertIn('Python',bold)
        self.assertNotIn('Alex Fixture',bold);self.assertNotIn('Technology Experience',bold)
        reader=PdfReader(self.root/app/row['files']['pdf']);fragments=[]
        for page in reader.pages:
            page.extract_text(visitor_text=lambda text,cm,tm,font,size:fragments.append((text,str((font or {}).get('/BaseFont','')))))
        self.assertTrue(any(tech['text'] in text and 'Bold' in font for text,font in fragments),fragments)
        self.assertTrue(any(duty['text'] in text and 'Bold' in font for text,font in fragments),fragments)
        self.assertTrue(any('Python' in text and 'Bold' in font for text,font in fragments))
        self.assertEqual([p.text for p in doc.paragraphs],(self.root/app/row['files']['txt']).read_text(encoding='utf-8').splitlines())
        self.assertEqual(self.store.profile('fixture'),before)
        compared=compare_resumes(self.store,self.root,app,base['id'],row['id'])
        self.assertEqual(compared['omitted'],[])
        self.assertIn('tech',[fact['fact_id'] for fact in compared['emphasized_skills']])
        self.assertIn('python-duty',[fact['fact_id'] for fact in compared['emphasized_skills']])
        from services.documents.uploads import save_resume
        from services.documents.uploaded_comparison import compare_uploaded_resume
        uploaded=save_resume(self.store,self.root,'fixture','base.txt','text/plain',original['txt'])
        uploaded_diff=compare_uploaded_resume(self.store,self.root,app,uploaded['revision_id'],row['id'])
        self.assertIn('tech',[fact['fact_id'] for fact in uploaded_diff['emphasized_skills']])
        prepare(self.store,job['id'],'fixture',self.root,tailor_resume=True)
        self.assertEqual(self.store.application(app)['document_version'],2)
        for ext,path in base['files'].items():self.assertEqual((self.root/app/path).read_bytes(),original[ext])

    def test_selection_does_not_emphasize_omitted_or_unmatched_facts(self):
        from sqlalchemy import update
        from packages.persistence import models as m
        self.store.save_profile(dict(PROFILE,facts=[*PROFILE['facts'],dict(PROFILE['facts'][0],id='other',kind='technology_experience',text='Maintained COBOL systems.')]))
        with fixture_server() as origin:discover(self.store,origin,'fixture')
        job=self.store.jobs()[0]
        with self.store.engine.begin() as c:c.execute(update(m.jobs).where(m.jobs.c.id==job['id']).values(skills=['Python']))
        app=prepare(self.store,job['id'],'fixture',self.root,tailor_resume=True,tailor_fact_ids=['other'])
        row=next(r for r in application_documents(self.store,app) if r['kind']=='resume')
        self.assertEqual(row['generation_context']['emphasized_fact_ids'],[])
        doc=Document(self.root/app/row['files']['docx'])
        self.assertFalse(any(r.bold for p in doc.paragraphs for r in p.runs))
        self.assertNotIn('Python',(self.root/app/row['files']['txt']).read_text())

    def test_historical_selection_still_validates_context_and_new_emphasis_cannot_lie(self):
        from sqlalchemy import update
        from packages.persistence import models as m
        from services.documents.comparison import ComparisonUnavailable
        with fixture_server() as origin:discover(self.store,origin,'fixture')
        job=self.store.jobs()[0];app=prepare(self.store,job['id'],'fixture',self.root)
        base=next(r for r in application_documents(self.store,app) if r['kind']=='resume')
        selected=[self.store.profile('fixture')['facts'][0]['fact_id']]
        prepare(self.store,job['id'],'fixture',self.root,tailor_resume=True,tailor_fact_ids=selected)
        row=next(r for r in application_documents(self.store,app) if r['kind']=='resume' and r['version']==2)
        with self.store.engine.begin() as c:c.execute(update(m.documents).where(m.documents.c.id==row['id']).values(generation_context={**row['generation_context'],'emphasized_fact_ids':['invented']}))
        with self.assertRaises(ComparisonUnavailable):compare_resumes(self.store,self.root,app,base['id'],row['id'])
        legacy={k:v for k,v in row['generation_context'].items() if k!='emphasized_fact_ids'}
        with self.store.engine.begin() as c:c.execute(update(m.documents).where(m.documents.c.id==row['id']).values(template='approved-fact-selection-headings-v1',generation_model='deterministic-fact-selection-v1',generation_context=legacy))
        compare_resumes(self.store,self.root,app,base['id'],row['id'])
        with self.store.engine.begin() as c:c.execute(update(m.documents).where(m.documents.c.id==row['id']).values(generation_context={**legacy,'available_fact_ids':[]}))
        with self.assertRaises(ComparisonUnavailable):compare_resumes(self.store,self.root,app,base['id'],row['id'])
