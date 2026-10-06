"""Deterministic claim checks, format corruption and concurrent font use."""
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import tempfile
import unittest

from docx import Document
from pypdf import PdfReader
from services.documents.errors import PreparationError
from services.documents.render import render_pdf
from services.documents.validation import validate_assertions, validate_formats


class DocumentValidationUnitTests(unittest.TestCase):
    def test_claim_evidence_records_protected_experience_and_education_fields(self):
        profile = {'name': 'Alex', 'email': 'alex@example.test', 'facts': [
            {'fact_id': 'work', 'kind': 'experience',
             'text': 'Managed Northstar Systems from 2020 to 2024; restored service 35 percent faster.',
             'verified': True, 'user_approved': True, 'confidence': 100, 'source': 'fixture'},
            {'fact_id': 'degree', 'kind': 'education',
             'text': 'Bachelor of Science, North Valley University, 2018.',
             'verified': True, 'user_approved': True, 'confidence': 100, 'source': 'fixture'},
        ]}
        evidence = validate_assertions(
            'resume',
            ['Alex', 'alex@example.test', profile['facts'][0]['text'], profile['facts'][1]['text']],
            profile, ['work', 'degree'])

        self.assertEqual(evidence['strategy'], 'DETERMINISTIC_EXACT_FACT_PRESERVATION')
        claims = {claim['fact_id']: claim for claim in evidence['claims']}
        self.assertEqual(set(claims), {'work', 'degree'})
        self.assertEqual(claims['work']['kind'], 'experience')
        self.assertEqual(claims['degree']['kind'], 'education')
        self.assertEqual(claims['work']['decision'], 'PRESERVED_EXACT')
        self.assertEqual(claims['work']['protected_fields'], [
            'FACT_TEXT', 'DATE_TOKENS', 'NUMERIC_TOKENS',
            'EMPLOYER_TEXT_WITHIN_EXPERIENCE_FACT'])
        self.assertEqual(claims['work']['protected_token_counts'], {'dates': 2, 'numbers': 3})
        self.assertNotIn('token_sha256', claims['work'])
        self.assertNotIn('fact_sha256', claims['work'])
        self.assertNotIn('text', claims['work'])

    def test_claim_mutations_to_metrics_dates_employers_or_degrees_are_rejected(self):
        text = 'Managed Northstar Systems from 2020 to 2024; restored service 35 percent faster.'
        profile = {'name': 'Alex', 'email': 'alex@example.test', 'facts': [
            {'fact_id': 'work', 'kind': 'experience', 'text': text,
             'verified': True, 'user_approved': True, 'confidence': 100, 'source': 'fixture'},
            {'fact_id': 'degree', 'kind': 'education',
             'text': 'Bachelor of Science, North Valley University, 2018.',
             'verified': True, 'user_approved': True, 'confidence': 100, 'source': 'fixture'},
        ]}
        mutations = [
            'Managed Northstar Systems from 2020 to 2025; restored service 35 percent faster.',
            'Managed Northstar Systems from 2020 to 2024; restored service 45 percent faster.',
            'Managed Southstar Systems from 2020 to 2024; restored service 35 percent faster.',
            'Bachelor of Arts, North Valley University, 2018.',
            'Bachelor of Science, South Valley University, 2018.',
        ]
        for mutation in mutations:
            with self.subTest(mutation=mutation), self.assertRaises(PreparationError) as error:
                if mutation.startswith('Bachelor'):
                    paragraphs = ['Alex', 'alex@example.test', text, mutation]
                else:
                    paragraphs = ['Alex', 'alex@example.test', mutation,
                                  profile['facts'][1]['text']]
                validate_assertions('resume', paragraphs, profile, ['work', 'degree'])
            self.assertEqual(error.exception.code, 'UNSUPPORTED_CLAIM')

    def test_assertions_reject_unapproved_ids_added_prose_and_changed_identity(self):
        profile={'name':'Alex','email':'alex@example.test','facts':[
            {'fact_id':'skill','text':'Python','verified':True,'user_approved':True,'confidence':100,'source':'fixture'},
            {'fact_id':'pending','text':'SQL','verified':True,'user_approved':False,'confidence':100,'source':'fixture'}]}
        for paragraphs,ids in [(['Alex','alex@example.test','SQL'],['pending']),
                (['Alex','alex@example.test','Python expert'],['skill']),
                (['Other','alex@example.test','Python'],['skill']),
                (['Alex','alex@example.test','Python','Python'],['skill','skill'])]:
            with self.subTest(paragraphs=paragraphs),self.assertRaises(PreparationError) as error:
                validate_assertions('resume',paragraphs,profile,ids)
            self.assertEqual(error.exception.code,'UNSUPPORTED_CLAIM')

    def test_parallel_pdf_rendering_and_format_corruption(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            texts=['Δοκιμή Система résumé '+str(i) for i in range(8)]
            def render(i):
                path=root/f'{i}.pdf';render_pdf(path,[texts[i]])
                return PdfReader(path).pages[0].extract_text().strip()
            with ThreadPoolExecutor(max_workers=4) as pool:
                self.assertEqual(list(pool.map(render,range(8))),texts)
            paths={'pdf':root/'0.pdf','docx':root/'test.docx','txt':root/'test.txt'}
            doc=Document();doc.add_paragraph(texts[0]);doc.save(paths['docx'])
            paths['txt'].write_text(texts[0]+'\n',encoding='utf-8')
            validate_formats(paths,[texts[0]])
            for fmt in ('txt','pdf','docx'):
                original=paths[fmt].read_bytes()
                if fmt=='txt':paths[fmt].write_text('Changed\n',encoding='utf-8')
                elif fmt=='pdf':render_pdf(paths[fmt],['Changed'])
                else:
                    doc=Document();doc.add_paragraph('Changed');doc.save(paths[fmt])
                with self.subTest(format=fmt),self.assertRaises(PreparationError) as error:
                    validate_formats(paths,[texts[0]])
                self.assertEqual(error.exception.code,'DOCUMENT_CONTENT_MISMATCH')
                paths[fmt].write_bytes(original)
