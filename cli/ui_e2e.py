"""Connected browser acceptance against the actual local Compose deployment."""
import argparse
from datetime import datetime, timezone
import hashlib
import io
import ipaddress
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import time
from urllib.parse import urlsplit

import httpx
from docx import Document
from playwright.sync_api import sync_playwright, expect
from reportlab.pdfgen import canvas

from cli.live_feedback import capture_redacted_screenshot
from cli.compose_recovery import collect_public_job_inventory
from cli import local_persistence_acceptance as persistence


_SAFE_OPERATION_ID = re.compile(r"[a-f0-9]{64}\Z")
_SAFE_OPERATION_ERROR = re.compile(r"[A-Z][A-Z0-9_]{0,63}\Z")
_SAFE_REQUEST_ID = re.compile(r"[A-Za-z0-9][A-Za-z0-9._:-]{0,127}\Z")
_APPLICATION_STATES = {
    "PREPARED", "REVIEW_REQUIRED", "AUTH_REQUIRED", "DRY_RUN_COMPLETE",
    "SUBMITTING", "SUBMISSION_UNCONFIRMED", "CONFIRMED",
}
_PREFLIGHT_RESULTS = {"PASS", "FAIL", "BLOCKED", "NOT_RUN", "UNKNOWN"}


def _terminal_operation_summary(*, operation_id, state, error_code, context):
    safe_state = state if state in {
        "COMPLETED", "FAILED", "CANCELLED", "BLOCKED", "DEAD_LETTERED", "UNKNOWN",
    } else "UNKNOWN"
    summary = {"kind": "APPLY_MOCK", "state": safe_state}
    if isinstance(operation_id, str) and _SAFE_OPERATION_ID.fullmatch(operation_id):
        summary["operation_id"] = operation_id
    if isinstance(error_code, str) and _SAFE_OPERATION_ERROR.fullmatch(error_code):
        summary["error_code"] = error_code
    if not isinstance(context, dict):
        return summary
    application_id = context.get("application_id")
    if isinstance(application_id, str) and _SAFE_OPERATION_ID.fullmatch(application_id):
        summary["application_id"] = application_id
    request_id = context.get("request_id")
    if isinstance(request_id, str) and _SAFE_REQUEST_ID.fullmatch(request_id):
        summary["request_id"] = request_id
    profile_version = context.get("expected_profile_version")
    if isinstance(profile_version, str) and _SAFE_OPERATION_ID.fullmatch(profile_version):
        summary["expected_profile_version"] = profile_version
    document_version = context.get("expected_document_version")
    if isinstance(document_version, int) and not isinstance(document_version, bool) and document_version >= 0:
        summary["expected_document_version"] = document_version
    application_state = context.get("application_state")
    if isinstance(application_state, str) and application_state in _APPLICATION_STATES:
        summary["application_state"] = application_state
    preflight_result = context.get("preflight_result")
    if isinstance(preflight_result, str) and preflight_result in _PREFLIGHT_RESULTS:
        summary["preflight_result"] = preflight_result
    for key in ("preflight_is_current", "submit"):
        value = context.get(key)
        if isinstance(value, bool):
            summary[key] = value
    return summary


class MockRunTerminalError(RuntimeError):
    """A terminal operation failure with only bounded public identifiers."""

    def __init__(self, *, operation_id, state, error_code, context=None):
        safe_state = state if state in {"FAILED", "CANCELLED", "BLOCKED", "DEAD_LETTERED"} else "UNKNOWN"
        super().__init__(f"mock run did not complete: {safe_state}")
        self.terminal_summary = _terminal_operation_summary(
            operation_id=operation_id, state=safe_state,
            error_code=error_code, context=context,
        )


class EvidenceJournal:
    """Atomically checkpoint only bounded, sanitized named UI evidence."""

    MAX_CHECKS = 80
    MAX_BYTES = 512 * 1024

    def __init__(self, path, binding):
        self.path = Path(path)
        self.binding = binding
        self.checks = []
        self.attempted = None
        self.failure = None
        self.overall_state = "IN_PROGRESS"
        self.remaining_state = "NOT_RUN"

    def begin(self, name):
        self.attempted = {"action": str(name)[:200], "result": "IN_PROGRESS"}
        self.persist()

    def complete(self, check):
        if len(self.checks) >= self.MAX_CHECKS:
            raise ValueError("evidence check limit exceeded")
        safe_check = {"action": str(check.get("action", "unnamed check"))[:300],
                      "result": "PASS"}
        if check.get("boundary"):
            safe_check["boundary"] = str(check["boundary"])[:500]
        self.checks.append(safe_check)
        self.attempted = None
        self.persist()

    def fail(self, error):
        stage = (self.attempted or {}).get("action", "Connected UI workflow")
        self.attempted = {"action": stage, "result": "FAIL"}
        self.failure = {"type": type(error).__name__, "stage": stage}
        if isinstance(error, MockRunTerminalError):
            self.failure["terminal_operation"] = error.terminal_summary
        self.overall_state = "FAIL"
        self.remaining_state = "NOT_RUN"
        self.persist()

    def finish(self, evidence):
        self.binding.update(evidence)
        self.overall_state = evidence.get("result", "PARTIAL")
        self.remaining_state = "NONE"
        self.attempted = None
        self.persist()

    def payload(self):
        checks = list(self.checks)
        if self.attempted:
            checks.append(dict(self.attempted))
        return {
            **self.binding,
            "overall_state": self.overall_state,
            "checks": checks,
            "remaining_state": self.remaining_state,
            **({"failure": self.failure} if self.failure else {}),
        }

    def persist(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        encoded = json.dumps(self.payload(), indent=2, sort_keys=True) + "\n"
        if len(encoded.encode("utf-8")) > self.MAX_BYTES:
            raise ValueError("evidence size limit exceeded")
        temporary = None
        try:
            with tempfile.NamedTemporaryFile(
                    mode="w", encoding="utf-8", dir=self.path.parent,
                    prefix=f".{self.path.name}.", suffix=".tmp", delete=False) as stream:
                temporary = Path(stream.name)
                stream.write(encoded)
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temporary, self.path)
        finally:
            if temporary is not None and temporary.exists():
                temporary.unlink()


def wait_for_mock_run(request, path, operation_id, *, failure_context=None, timeout_seconds=210,
                      clock=time.monotonic, sleep=time.sleep):
    """Wait for this reviewed intent, without resubmitting or hiding worker failure."""
    deadline = clock() + timeout_seconds
    while True:
        result = request(path)
        if result.get("operation_id") != operation_id:
            raise RuntimeError("mock run identity changed")
        state_value = result.get("state")
        state = state_value if isinstance(state_value, str) else "UNKNOWN"
        if state == "COMPLETED":
            return result
        if state not in {"QUEUED", "RUNNING"}:
            # Do not publish arbitrary worker diagnostics or untrusted state text.
            safe_state = state if state in {"FAILED", "CANCELLED", "BLOCKED", "DEAD_LETTERED"} else "UNKNOWN"
            raise MockRunTerminalError(
                operation_id=operation_id,
                state=safe_state,
                error_code=result.get("error_code"),
                context=failure_context,
            )
        remaining = deadline - clock()
        if remaining <= 0:
            raise TimeoutError(f"mock run completion timed out: {state}")
        sleep(min(1, remaining))


def validate_onboarding_preflight(settings, health, readiness):
    if (settings.get("operation_mode") != "FIXTURE"
            or settings.get("external_submission_enabled") is not False
            or health.get("operation_mode") != "FIXTURE"
            or health.get("external_submission_enabled") is not False
            or readiness.get("status") != "ready"):
        raise RuntimeError("onboarding requires ready FIXTURE mode with external submission disabled")


def validate_onboarding_target(base_url):
    parsed = urlsplit(base_url)
    hostname = (parsed.hostname or "").lower()
    loopback = hostname == "localhost"
    try:
        loopback = loopback or ipaddress.ip_address(hostname).is_loopback
    except ValueError:
        pass
    if parsed.scheme != "http" or not loopback or parsed.username or parsed.password:
        raise RuntimeError("onboarding target must be an HTTP loopback Compose service")


def assert_onboarding_profile_facts(profile, pending_fact_ids):
    facts = profile["facts"]
    by_id = {row["fact_id"]: row for row in facts}
    pending_fact_ids = set(pending_fact_ids)
    assert len(by_id) == len(facts), "duplicate profile fact IDs"
    approved = [row for row in facts if row["verified"] and row["user_approved"]]
    if ({row["text"] for row in approved} != {"Python", "SQL"}
            or len(approved) != 2
            or any(row["kind"] != "skill" for row in approved)
            or any(row["fact_id"] in pending_fact_ids for row in approved)):
        raise AssertionError("unexpected approved facts in onboarding profile")
    pending = [row for row in facts if row["fact_id"] in pending_fact_ids]
    if (len(pending) != len(pending_fact_ids)
            or any(row["verified"] or row["user_approved"] for row in pending)
            or set(by_id) != pending_fact_ids | {row["fact_id"] for row in approved}):
        raise AssertionError("onboarding profile facts differ from reviewed fact set")
    return {row["text"]: row["fact_id"] for row in approved}


def assert_no_pending_facts_in_documents(documents, pending_fact_ids, approved_fact_ids):
    pending_fact_ids = set(pending_fact_ids)
    approved_fact_ids = set(approved_fact_ids)
    for document in documents:
        used_ids = set(document.get("fact_ids", []))
        facts_used = document.get("facts_used", [])
        used_ids.update(row.get("fact_id") for row in facts_used)
        if (used_ids & pending_fact_ids or used_ids - approved_fact_ids
                or any(not row.get("verified") or not row.get("user_approved")
                       for row in facts_used)):
            raise AssertionError("pending or unsupported facts reached generated documents")


def assert_search_scoped_discovery(request_payload, operation, jobs, *, profile_id,
                                   search_profile_id, expected_titles):
    """Check queued fixture discovery and durable readback for a selected saved search."""
    if request_payload.get("profile_id") != profile_id:
        raise AssertionError("discovery request changed the active profile")
    if request_payload.get("search_profile_id") != search_profile_id:
        raise AssertionError("discovery request omitted or changed the selected saved search")
    if operation.get("entity_id") != profile_id:
        raise AssertionError("discovery operation changed the active profile")
    params = operation.get("params") or {}
    if params.get("profile_id", profile_id) != profile_id:
        raise AssertionError("discovery operation parameters changed the active profile")
    if params.get("search_profile_id") != search_profile_id:
        raise AssertionError("discovery operation did not retain the selected saved search")
    if any(row.get("profile_id", profile_id) != profile_id for row in jobs):
        raise AssertionError("search-filtered readback included another profile's job")
    titles = [row.get("title") for row in jobs]
    if set(titles) != set(expected_titles) or len(titles) != len(expected_titles):
        raise AssertionError("search-filtered job readback differs from the reviewed criteria")


def assert_discovery_request_operation_pair(request_payload, operation, *, profile_id,
                                            search_profile_id, expected_operation_id):
    """Tie a submitted request to its operation; entity_id owns the profile scope."""
    params = operation.get("params") if isinstance(operation, dict) else None
    if (not isinstance(request_payload, dict) or not isinstance(params, dict)
            or not isinstance(request_payload.get("request_id"), str)
            or not request_payload["request_id"]
            or not isinstance(operation.get("id"), str) or not operation["id"]
            or operation["id"] != expected_operation_id
            or request_payload.get("profile_id") != profile_id
            or operation.get("entity_id") != profile_id
            or request_payload.get("search_profile_id") != search_profile_id
            or params.get("search_profile_id") != search_profile_id
            or params.get("request_id") != request_payload["request_id"]):
        raise AssertionError("discovery request and persisted operation do not match")
    return operation["id"], request_payload["request_id"]


def assert_supported_fixture_search(search, *, profile_id, include_titles, remote):
    """Require a Sources-tab discovery search, never a Jobs filter snapshot."""
    criteria = search.get("criteria") if isinstance(search, dict) else None
    if not isinstance(criteria, dict):
        raise AssertionError("saved fixture search has no criteria")
    if search.get("profile_id") != profile_id or search.get("enabled") is not True:
        raise AssertionError("fixture search is disabled or belongs to another profile")
    if search.get("source_ids") != []:
        raise AssertionError("fixture discovery requires the all-sources search scope")
    if (criteria.get("include_titles") != include_titles or criteria.get("remote") != remote
            or criteria.get("employment_types") != []):
        raise AssertionError("saved fixture search differs from the reviewed Sources criteria")
    if "job_filters" in criteria:
        raise AssertionError("Jobs-tab filters cannot be used as provider discovery criteria")
    scheduler = criteria.get("scheduler")
    if not isinstance(scheduler, dict) or scheduler.get("enabled") is not False:
        raise AssertionError("fixture search must keep its background scheduler disabled")


def assert_reprepared_profile_pins(*, previous_profile_version, old_application_profile_version,
                                   current_profile_version, resumed_application_profile_version):
    """Compare opaque profile-version identities by equality, not lexical order."""
    versions = (previous_profile_version, old_application_profile_version,
                current_profile_version, resumed_application_profile_version)
    if any(not isinstance(version, str) or not version for version in versions):
        raise AssertionError("reprepare profile versions must be nonempty identities")
    if previous_profile_version == current_profile_version:
        raise AssertionError("reviewed profile edit did not produce a new profile version")
    if old_application_profile_version != previous_profile_version:
        raise AssertionError("application did not retain its old profile-version pin before reprepare")
    if resumed_application_profile_version != current_profile_version:
        raise AssertionError("reprepared application is not pinned to the approved current profile")


def assert_blocked_mock_operation(operation, *, application_id):
    """Validate the legacy public operation record used by onboarding recovery."""
    assert operation.get("kind") == "APPLY_MOCK", operation
    assert operation.get("entity_id") == application_id, operation
    assert operation.get("state") == "BLOCKED", operation


def assert_previous_review_operation(request, *, onboarding, operation_id, application_id):
    """Keep the onboarding-only blocked receipt check out of the demo flow."""
    if not onboarding:
        return
    assert operation_id, "onboarding recovery requires the prior blocked operation"
    operation = request("/operations/" + operation_id)
    assert_blocked_mock_operation(operation, application_id=application_id)


def is_allowed_onboarding_application_write(method, path):
    """Allow only the explicit local application operations used by onboarding."""
    if method != "POST":
        return False
    if path == "/applications":
        return True
    parts = path.strip("/").split("/")
    if (len(parts) == 5 and parts[0] == "profiles" and bool(parts[1])
            and parts[2] == "applications" and bool(parts[3])
            and parts[4] in {"reprepare", "mock-run"}):
        return True
    comparison_path = (len(parts) == 6 and parts[0] == "profiles" and bool(parts[1])
                       and parts[2] == "applications" and bool(parts[3])
                       and parts[4:] == ["uploaded-resume-comparison", "operations"])
    comparison_cancel_path = (len(parts) == 8 and parts[0] == "profiles" and bool(parts[1])
                              and parts[2] == "applications" and bool(parts[3])
                              and parts[4:6] == ["uploaded-resume-comparison", "operations"]
                              and bool(parts[6]) and parts[7] == "cancel")
    return comparison_path or comparison_cancel_path


def is_allowed_onboarding_request(method, url, allowed_origin):
    """Enforce same-origin, no employer submission, and exact local write routes."""
    parsed = urlsplit(url)
    path = parsed.path.lower()
    blocked_path = (path.rstrip("/").endswith(("/apply", "/submit"))
                    or path in {"/apply", "/submit"}
                    or "/external-application" in path)
    application_write = method != "GET" and "/applications/" in path
    return (parsed.netloc == allowed_origin and not blocked_path
            and (not application_write
                 or is_allowed_onboarding_application_write(method, path)))


def sanitized_onboarding_request(method, url):
    """Return method and a query-free route path with profile/application IDs hidden."""
    parts = urlsplit(url).path.strip("/").split("/")
    if len(parts) >= 4 and parts[0].lower() == "profiles" and parts[2].lower() == "applications":
        parts[1] = "<profile>"
        parts[3] = "<application>"
    path = "/" + "/".join(parts) if parts and parts != [""] else "/"
    return {"method": method.upper(), "path": path}


def expect_visible_with_onboarding_diagnostics(locator, blocked_requests, **kwargs):
    """Keep local browser failures actionable without recording request payloads."""
    try:
        expect(locator).to_be_visible(**kwargs)
    except AssertionError as error:
        if blocked_requests:
            raise AssertionError(
                f"{error}\nBlocked onboarding requests (method/path only): {blocked_requests}"
            ) from error
        raise


def answer_bank_confirmation_locator(answer_group, question_number):
    """Locate the confirmation div by its aria-label (it is not a region role)."""
    return answer_group.locator(
        f'[aria-label="Answer Bank confirmation for question {question_number}"]'
    )


def _onboarding_fixture_pdf():
    stream = io.BytesIO()
    pdf = canvas.Canvas(stream)
    text = pdf.beginText(40, 780)
    for line in (
        "Summary", "Fictional onboarding summary for extraction review only.",
    ):
        text.textLine(line)
    pdf.drawText(text)
    pdf.save()
    return stream.getvalue()


def _onboarding_preflight(base_url):
    validate_onboarding_target(base_url)
    with httpx.Client(base_url=base_url, timeout=5, trust_env=False) as client:
        settings = client.get("/settings")
        settings.raise_for_status()
        health = client.get("/health")
        health.raise_for_status()
        readiness = client.get("/readiness")
        readiness.raise_for_status()
    validate_onboarding_preflight(settings.json(), health.json(), readiness.json())


def wait_for_ui_action_idle(page, button_name):
    """Wait for the shared UI busy gate to release a form action."""
    expect(page.get_by_role("button", name=button_name, exact=True)).to_be_enabled()


def wait_for_fixture_collection_ui_idle(page):
    """Wait until fixture collection, its UI refresh, and the shared busy gate finish."""
    expect(page.get_by_role("status").filter(
        has_text="Local fixture jobs collected for your saved profile. No employer was contacted."
    )).to_be_visible()
    # The collection action navigates back to Jobs before action() completes its
    # final refresh/readback. This stable Jobs control remains mounted through
    # that transition and is disabled by the same busy gate as the Sources form.
    wait_for_ui_action_idle(page, "Refresh matches")


def _source_binding():
    def git_value(*args):
        result = subprocess.run(
            ["git", *args], capture_output=True, text=True, timeout=5, check=True)
        return result.stdout.strip()

    try:
        return {"commit": git_value("rev-parse", "HEAD"),
                "tree": git_value("rev-parse", "HEAD^{tree}"),
                "working_tree_diff_sha256": hashlib.sha256(
                    git_value("diff", "--binary", "HEAD").encode("utf-8")
                ).hexdigest(),
                "working_tree_status_sha256": hashlib.sha256(
                    git_value("status", "--porcelain=v1", "--untracked-files=all").encode("utf-8")
                ).hexdigest()}
    except (OSError, subprocess.SubprocessError):
        return {"commit": None, "tree": None, "state": "UNAVAILABLE"}


def run(base_url, evidence_path, onboarding=False):
    if onboarding:
        # This read-only gate runs before output creation, Chromium launch, or any UI write.
        _onboarding_preflight(base_url)
    artifacts = Path("outputs/ui-e2e") / str(urlsplit(base_url).port)
    artifacts.mkdir(parents=True, exist_ok=True)
    parsed_target = urlsplit(base_url)
    journal = EvidenceJournal(evidence_path, {
        "schema_version": 1,
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "mode": "ONBOARDING_FIXTURE" if onboarding else "FICTIONAL_DEMO",
        "target_origin": f"{parsed_target.scheme}://{parsed_target.netloc}",
        "source": _source_binding(),
    })

    class CheckpointActions(list):
        def append(self, value):
            if value.get("result") == "PASS":
                journal.complete(value)
            super().append(value)

    actions = CheckpointActions()
    job_inventory = None
    with sync_playwright() as browser_api:
        browser = browser_api.chromium.launch(headless=True)
        context = browser.new_context(viewport={"width": 1440, "height": 1050}, accept_downloads=True)
        unsafe_requests = []
        if onboarding:
            allowed_origin = urlsplit(base_url).netloc

            def local_fixture_route(route):
                current = route.request
                if not is_allowed_onboarding_request(current.method, current.url, allowed_origin):
                    unsafe_requests.append(sanitized_onboarding_request(current.method, current.url))
                    route.abort()
                    return
                route.continue_()

            context.route("**/*", local_fixture_route)
        page = context.new_page()
        page.set_default_timeout(45000)
        failures = []
        page.on("pageerror", lambda failure: failures.append(str(failure)))

        def navigation(name):
            page.get_by_role("navigation").get_by_role("button", name=name).click()

        def request(path):
            response = context.request.get(base_url + path)
            assert response.ok, (path, response.status)
            return response.json()

        def download_document(document_id, format):
            path = f"/profiles/{profile_id}/documents/{document_id}/{format}"
            response = context.request.get(base_url + path)
            assert response.ok, (path, response.status)
            return response.body()

        def save_profile():
            save_button = page.get_by_role("button", name="Save profile", exact=True)
            save_button.click()
            expect(page.get_by_role("status").filter(has_text="Profile saved")).to_be_visible()
            # The notice is set before action() finishes its refresh and jobs readback.
            # Wait for the actual shared busy gate to clear before continuing to another form.
            wait_for_ui_action_idle(page, "Save profile")

        def fact_editor(text):
            for fieldset in page.locator("fieldset.fact").all():
                value = fieldset.get_by_role("textbox", name="Fact text", exact=True)
                if value.count() and value.input_value() == text:
                    return fieldset
            raise AssertionError("Expected exact profile fact text was not found: " + text)

        def prepare():
            if page.get_by_label("Order approved skills for this job", exact=True).is_checked():
                preview = page.get_by_role("region", name="Resume keyword alignment preview", exact=True)
                preview.get_by_role("button", name="Preview keyword alignment", exact=True).click()
                expect(preview.get_by_role("status")).to_contain_text("Read-only preview")
            page.get_by_role("button", name="Prepare documents", exact=True).click()
            expect(page.get_by_role("heading", name="Application details", exact=True)).to_be_visible()
            expect(page.get_by_role("button", name="Run dry-run", exact=True)).to_be_enabled()

        def select_data_job():
            navigation("Jobs")
            page.get_by_label("Search jobs").fill("data")
            if onboarding:
                page.get_by_label("Job title", exact=True).fill("")
                page.get_by_label("Work location", exact=True).select_option("")
            card = page.get_by_role("button").filter(has=page.get_by_role("heading", name="Data Engineer", exact=True))
            expect(card).to_have_count(1)
            card.click()
            expect(page.get_by_role("heading", name="Why this match?", exact=True)).to_be_visible()

        try:
            page.goto(base_url, wait_until="domcontentloaded")
            expect(page.get_by_role("heading", name="Find your next opportunity.")).to_be_visible()
            pending_fact_ids = set()
            discovery_search_proofs = []
            if onboarding:
                page.get_by_role("button", name="New profile", exact=True).click()
                page.get_by_label("Full name", exact=True).fill("Synthetic Onboarding Applicant")
                page.get_by_label("Email", exact=True).fill("onboarding@example.invalid")
                with page.expect_response(lambda response:
                        response.request.method == "PUT"
                        and urlsplit(response.url).path == "/profiles") as profile_saved:
                    save_profile()
                assert profile_saved.value.status == 200, profile_saved.value.text()
                profile_id = profile_saved.value.json()["id"]
                expect(page.get_by_label("Active profile", exact=True)).to_have_value(profile_id)
                assert profile_id and not profile_id.startswith("demo-")
                profile = request("/profiles/" + profile_id)
                assert profile["facts"] == []
                actions.append({"action": "UI created a new synthetic profile without replacing existing profiles", "result": "PASS"})

                page.get_by_label("Choose a resume (up to 10 MiB)").set_input_files({
                    "name": "synthetic-onboarding.pdf", "mimeType": "application/pdf",
                    "buffer": _onboarding_fixture_pdf(),
                })
                page.get_by_label("Resume label", exact=True).fill("Synthetic onboarding fixture")
                page.get_by_role("button", name="Store resume", exact=True).click()
                resume_row = page.locator("[data-resume-version='1']")
                expect(resume_row).to_be_visible()
                resume_row.get_by_role("button", name="Extract and review", exact=True).click()
                review = page.get_by_role("region", name="Resume extraction review")
                expect(review).to_be_visible()
                suggestions = review.locator('label.fact-suggestion:has(strong) input[type="checkbox"]')
                assert suggestions.count() > 0
                for suggestion in suggestions.all():
                    expect(suggestion).not_to_be_checked()
                    suggestion.check()
                expect(review.get_by_role("button", name="Add selected to profile editor", exact=True)).to_be_enabled()
                review.get_by_role("button", name="Add selected to profile editor", exact=True).click()
                imported_rows = page.locator("fieldset.fact")
                expect(imported_rows).not_to_have_count(0, timeout=10000)
                for imported_fact in imported_rows.all():
                    expect(imported_fact).to_contain_text("Pending")
                    expect(imported_fact.get_by_label("I verified this fact", exact=True)).not_to_be_checked()
                    expect(imported_fact.get_by_label("I approve its use in documents", exact=True)).not_to_be_checked()
                save_profile()
                imported_profile = request("/profiles/" + profile_id)
                pending_fact_ids = {row["fact_id"] for row in imported_profile["facts"]}
                assert pending_fact_ids
                assert all(not row["verified"] and not row["user_approved"]
                           for row in imported_profile["facts"])
                actions.append({"action": "PDF imported through UI and extraction suggestions remain unverified and unapproved after readback", "result": "PASS"})

                for fact_text in ("Python", "SQL"):
                    page.get_by_role("button", name="Add fact", exact=True).click()
                    new_fact_editor = page.locator("fieldset.fact").last
                    new_fact_editor.get_by_role("combobox").select_option("skill")
                    new_fact_editor.get_by_label("Confidence (0–100)", exact=True).fill("100")
                    new_fact_editor.get_by_label("Fact text", exact=True).fill(fact_text)
                    new_fact_editor.get_by_label("Evidence / source", exact=True).fill("Explicitly reviewed synthetic onboarding fact")
                    new_fact_editor.get_by_label("I verified this fact", exact=True).check()
                    new_fact_editor.get_by_label("I approve its use in documents", exact=True).check()
                for question, answer_text in (
                    ("Country", "United States"),
                    ("Remote role", "yes"),
                    ("Approved fixture answer", "Controlled fictional test application."),
                ):
                    page.get_by_role("button", name="Add answer", exact=True).click()
                    answer_editor = page.locator("fieldset.fact").last
                    answer_editor.get_by_role("textbox", name="Question this answer addresses", exact=True).fill(question)
                    answer_editor.get_by_role("textbox", name="Answer text", exact=True).fill(answer_text)
                    answer_editor.get_by_label("Provenance", exact=True).fill("Explicit synthetic mock-form review")
                    answer_editor.get_by_label("Confidence (0–100)", exact=True).fill("100")
                    if question != "Approved fixture answer":
                        answer_editor.get_by_label("I reviewed and verified this answer", exact=True).check()
                        answer_editor.get_by_label("I approve exact reuse for this question", exact=True).check()
                save_profile()
                profile = request("/profiles/" + profile_id)
                assert len(profile["answers"]) == 3
                answer_by_question = {row["question"]: row for row in profile["answers"]}
                assert {"Country", "Remote role", "Approved fixture answer"} == set(answer_by_question)
                assert all(answer_by_question[question]["verified"]
                           and answer_by_question[question]["user_approved"]
                           and answer_by_question[question]["confidence"] == 100
                           for question in ("Country", "Remote role"))
                assert not answer_by_question["Approved fixture answer"]["verified"]
                assert not answer_by_question["Approved fixture answer"]["user_approved"]
                assert not {row["id"] for row in profile["answers"]} & {"country", "remote", "fixture_question"}
                actions.append({"action": "UI saves two explicitly approved synthetic answers and keeps the required mock-form answer unreviewed for later question review", "result": "PASS"})
                skill_fact_ids = assert_onboarding_profile_facts(profile, pending_fact_ids)
                python_fact_id = skill_fact_ids["Python"]
                sql_fact_id = skill_fact_ids["SQL"]
                actions.append({"action": "Only explicitly entered and reviewed Python and SQL facts are verified and approved", "result": "PASS"})

                def save_sources_fixture_search(name, title):
                    navigation("Sources")
                    page.get_by_label("Search name", exact=True).fill(name)
                    page.get_by_label("Remote preference").select_option("REQUIRED")
                    page.get_by_label("Employment type").select_option("")
                    page.get_by_label("Include titles", exact=True).fill(title)
                    with page.expect_response(lambda candidate:
                            candidate.request.method == "PUT"
                            and urlsplit(candidate.url).path == "/search-profiles") as saved:
                        journal.begin("Save selected Sources search")
                        page.get_by_role("button", name="Save search", exact=True).click()
                    assert saved.value.status == 200, saved.value.text()
                    expect(page.get_by_role("status").filter(
                        has_text="Search saved. Its schedule is off")).to_be_visible()
                    search = next(row for row in request(
                        "/search-profiles?profile_id=" + profile_id)["search_profiles"]
                        if row["name"] == name)
                    assert_supported_fixture_search(
                        search, profile_id=profile_id, include_titles=[title], remote="REQUIRED")
                    navigation("Jobs")
                    page.get_by_label("Saved search profile", exact=True).last.select_option(search["id"])
                    expect(page.get_by_label("Saved search profile", exact=True).last).to_have_value(search["id"])
                    return search

                def collect_sources_fixture_search(search, expected_titles):
                    # Search selection happens in Jobs; the collection control is in Sources.
                    navigation("Sources")
                    with page.expect_response(lambda response:
                            response.request.method == "POST"
                            and urlsplit(response.url).path == "/discovery") as discovery_response:
                        page.get_by_role("button", name="Collect local fixture jobs", exact=True).click()
                    response = discovery_response.value
                    assert response.status == 202, response.text()
                    discovery_request = response.request.post_data_json
                    assert discovery_request["search_profile_id"] == search["id"]
                    operation_id = response.json()["operation_id"]
                    deadline = time.monotonic() + 120
                    while time.monotonic() < deadline:
                        operation = request("/operations/" + operation_id)
                        if operation["state"] in ("COMPLETED", "FAILED", "BLOCKED", "CANCELLED"):
                            break
                        time.sleep(0.5)
                    assert operation["state"] == "COMPLETED", operation
                    operation_id, request_id = assert_discovery_request_operation_pair(
                        discovery_request, operation, profile_id=profile_id,
                        search_profile_id=search["id"], expected_operation_id=operation_id)
                    jobs = request("/jobs?profile_id=" + profile_id
                                   + "&search_profile_id=" + search["id"])["jobs"]
                    assert_search_scoped_discovery(
                        discovery_request, operation, jobs,
                        profile_id=profile_id, search_profile_id=search["id"],
                        expected_titles=set(expected_titles))
                    # Operation completion and filtered readback can precede the
                    # UI action's own refresh/loadJobs and busy-gate release.
                    wait_for_fixture_collection_ui_idle(page)
                    discovery_search_proofs.append({
                        "search_profile_id": search["id"],
                        "operation_id": operation_id,
                        "request_id": request_id,
                        "profile_id": profile_id,
                        "filtered_titles": [row["title"] for row in jobs],
                        "result": "PASS",
                    })
                    return operation_id, discovery_request["request_id"], jobs

                # Create searches through the product's supported Sources form.
                # Jobs-tab filter snapshots are intentionally disabled for scheduling
                # and cannot be resumed as provider discovery criteria.
                nonmatching_search = save_sources_fixture_search(
                    "Synthetic remote IT onboarding search", "IT Support Specialist")
                discovery_operation_id, discovery_request_id, nonmatching_jobs = (
                    collect_sources_fixture_search(nonmatching_search, set()))
                assert {row["title"] for row in request(
                    "/jobs?profile_id=" + profile_id)["jobs"]} >= {"Python Developer", "Data Engineer"}
                actions.append({"action": "Created a remote IT title search in the Sources UI; actual fixture discovery completed and the selected same-profile filter returned no matching jobs", "result": "PASS", "boundary": "Sources search form → Jobs saved-search selection → queued /discovery → Compose source worker → selected-search jobs API"})

                matching_search = save_sources_fixture_search(
                    "Synthetic remote data engineering onboarding search", "Data Engineer")
                discovery_operation_id, discovery_request_id, matching_jobs = (
                    collect_sources_fixture_search(matching_search, {"Data Engineer"}))
                discovery_job_ids = {row["id"] for row in matching_jobs}
                assert discovery_job_ids, "matching saved search returned no fixture jobs"
                actions.append({"action": "Selected a matching remote Data Engineer search; actual fixture worker completed and only the matching job appeared in same-profile filtered readback", "result": "PASS", "boundary": "Jobs UI → Sources UI → queued /discovery → Compose source worker → selected-search jobs API"})
                history = request("/profiles/" + profile_id + "/search-history")
                assert history["profile_id"] == profile_id
                history_item = next(item for item in history["items"]
                                    if item["id"].startswith(discovery_operation_id + ":"))
                assert history_item["status"] == "COMPLETED"
                assert history_item["profile_id"] == profile_id
                assert history_item["source"] == "mock-fixture"
                assert history_item["results"] and history_item["results"] > 0
                persisted_operation = request("/operations/" + discovery_operation_id)
                assert history_item["id"] == (discovery_operation_id + ":"
                                               + str(persisted_operation["attempt_count"]))
                assert_discovery_request_operation_pair(
                    {"request_id": discovery_request_id, "profile_id": profile_id,
                     "search_profile_id": matching_search["id"]},
                    persisted_operation, profile_id=profile_id,
                    search_profile_id=matching_search["id"],
                    expected_operation_id=discovery_operation_id)
                search_after_discovery = next(row for row in request(
                    "/search-profiles?profile_id=" + profile_id)["search_profiles"]
                    if row["id"] == matching_search["id"])
                assert search_after_discovery["criteria"]["include_titles"] == ["Data Engineer"]
                assert search_after_discovery["criteria"]["remote"] == "REQUIRED"
                assert "job_filters" not in search_after_discovery["criteria"]
                assert search_after_discovery["source_ids"] == []
            else:
                page.get_by_role("button", name="Load fictional demo", exact=True).click()
                expect(page.get_by_role("status").filter(has_text="Fictional demo loaded")).to_be_visible()
                profile_id = page.get_by_label("Active profile").input_value()
                assert profile_id.startswith("demo-")
                profile = request("/profiles/" + profile_id)
                python_fact_id, sql_fact_id = "python", "sql"
            select_data_job()
            expect(page.locator(".big-score")).to_have_text("100% · Exceptional")
            actions.append({"action": "Explicit fictional demo, connected search, match explanation and raw provenance", "result": "PASS"})

            navigation("Profile")
            fact_editor("SQL").get_by_label("I approve its use in documents", exact=True).uncheck()
            save_profile()
            select_data_job()
            expect(page.locator(".big-score")).to_have_text("50% · Weak")
            prepare()
            application = request("/applications?profile_id=" + profile_id)["applications"][0]
            identifier = application["id"]
            if onboarding:
                assert application["job_id"] in discovery_job_ids, "selected application job was not returned by this profile's fixture discovery"
            assert application["fact_ids"] == [python_fact_id]
            first_documents = request("/profiles/" + profile_id + "/applications/" + identifier + "/documents")["documents"]
            first_resume = next(row for row in first_documents if row["kind"] == "resume")
            assert first_resume['created_at'] and first_resume['generation_model'] == 'deterministic-fact-copy-v2'
            assert [fact['text'] for fact in first_resume['facts_used']] == ['Python']
            provenance = page.get_by_role('region', name='resume version 1', exact=True)
            provenance.get_by_text('Generation details and facts', exact=True).click()
            expect(provenance.get_by_text('Generation method: deterministic-fact-copy-v2', exact=True)).to_be_visible()
            with page.expect_download() as text_download:
                provenance.get_by_role('link', name='Download TXT', exact=True).click()
            text_bytes = Path(text_download.value.path()).read_bytes()
            assert hashlib.sha256(text_bytes).hexdigest() == first_resume['sha256']['txt']
            original = download_document(first_resume["id"], "docx")
            text = "\n".join(row.text for row in Document(io.BytesIO(original)).paragraphs)
            assert "Python" in text and "SQL" not in text
            assert first_resume['template']=='verified-headings-v1'
            assert [p.text for p in Document(io.BytesIO(original)).paragraphs if p.style.name=='Heading 1']==['Skills']
            actions.append({'action':'ATS resume templates expose conventional headings in the browser-downloaded DOCX while preserving approved wording','result':'PASS'})
            assert text_bytes.decode('utf-8').strip() == text
            actions.append({'action':'Original approved facts, generation metadata and checksum-verified TXT download visible through the browser', 'result':'PASS'})
            actions.append({"action": "Pending fact excluded from matching and downloaded factual DOCX", "result": "PASS"})

            navigation("Profile")
            fact_editor("SQL").get_by_label("I approve its use in documents", exact=True).check()
            save_profile()
            select_data_job()
            expect(page.locator(".big-score")).to_have_text("100% · Exceptional")
            prepare()
            revised = request("/profiles/" + profile_id + "/applications/" + identifier)
            assert revised["document_version"] == 2 and set(revised["fact_ids"]) == {python_fact_id, sql_fact_id}
            documents = request("/profiles/" + profile_id + "/applications/" + identifier + "/documents")["documents"]
            assert len(documents) == 4
            old = download_document(first_resume["id"], "docx")
            assert old == original
            with page.expect_download() as download:
                page.get_by_role("link", name="Download PDF", exact=True).first.click()
            pdf_path = artifacts / "fictional-document.pdf"
            download.value.save_as(pdf_path)
            assert pdf_path.read_bytes().startswith(b"%PDF")
            assert hashlib.sha256(pdf_path.read_bytes()).hexdigest() in {doc["sha256"]["pdf"] for doc in documents}
            actions.append({"action": "Approved profile revision generates new documents, preserves original bytes, and PDF downloads through UI", "result": "PASS"})

            select_data_job()
            page.get_by_label('Order approved skills for this job',exact=True).check()
            prepare()
            tailored=next(row for row in request('/profiles/'+profile_id+'/applications/'+identifier+'/documents')['documents'] if row['version']==3 and row['kind']=='resume')
            assert tailored['generation_model']=='deterministic-target-relevant-bullets-summary-skill-responsibility-order-v4'
            assert tailored['template']=='approved-target-relevant-bullets-summary-skill-responsibility-and-bullet-emphasis-v5'
            assert tailored['generation_context']['emphasized_fact_ids']
            expect(page.get_by_role('region',name='resume version 3',exact=True).get_by_text('Job-matched approved facts in bold: 2.',exact=False)).to_be_visible()
            emphasized_doc=Document(io.BytesIO(download_document(tailored['id'],'docx')))
            assert any(p.text=='Python' and all(r.bold for r in p.runs) for p in emphasized_doc.paragraphs)
            actions.append({'action':'Job-matched technology emphasis reaches worker, DOCX download and visible provenance','result':'PASS'})
            assert tailored['generation_context']['mode']=='APPROVED_SKILL_ORDER'
            assert set(tailored['fact_ids'])=={python_fact_id,sql_fact_id}
            expect(page.get_by_role('region',name='resume version 3',exact=True).get_by_text('Tailored resume: skill order',exact=False)).to_be_visible()
            actions.append({'action':'Explicit skill-order opt-in reaches the document worker and preserves approved facts with target snapshot and separate revision','result':'PASS'})

            comparison=page.get_by_role('region',name='Compare base and tailored resumes',exact=True)
            comparison.get_by_label('Base resume',exact=True).select_option(label='v2')
            comparison.get_by_role('button',name='Compare selected resumes',exact=True).click()
            result=comparison.get_by_role('region',name='Resume comparison result',exact=True)
            expect(result.get_by_role('region',name='Content omitted',exact=True)).to_contain_text('No content omitted.')
            expect(result.get_by_role('region',name='Summary changes',exact=True)).to_contain_text('No changes.')
            expect(result.get_by_role('heading',name='Generated base · v2',exact=True)).to_be_visible()
            expect(result.get_by_role('heading',name='Tailored resume · v3',exact=True)).to_be_visible()
            actions.append({'action':'Immutable generated base and tailored resume comparison shows recorded versions and unchanged facts through the UI','result':'PASS'})

            profile_before_comparison=request('/profiles/'+profile_id)
            uploaded_original_text=b'Summary\nOriginal synthetic summary\nSkills\nSQL; Python\nResponsibilities\nMaintained synthetic servers\n'
            uploaded_response=context.request.post(base_url+'/profiles/'+profile_id+'/resumes', data=uploaded_original_text,
                headers={'X-Filename':'comparison-base.txt','Content-Type':'text/plain'})
            assert uploaded_response.status==201, uploaded_response.text()
            uploaded_original=uploaded_response.json()
            comparison.get_by_role('button',name='Refresh base resumes',exact=True).click()
            comparison.get_by_label('Base resume',exact=True).select_option('u:'+uploaded_original['revision_id'])
            comparison.get_by_role('button',name='Compare selected resumes',exact=True).click()
            expect_visible_with_onboarding_diagnostics(
                result.get_by_role('heading',name='Uploaded base · v1',exact=True), unsafe_requests)
            expect(result.get_by_role('region',name='Summary changes',exact=True)).to_contain_text('Original synthetic summary')
            expect(result.get_by_role('region',name='Content omitted',exact=True)).to_contain_text('Maintained synthetic servers')
            expect(result.get_by_role('region',name='Original extraction provenance',exact=True)).to_contain_text('Native extracted text')
            with page.expect_download() as original_download:
                result.get_by_role('link',name='Download compared original',exact=True).click()
            assert Path(original_download.value.path()).read_bytes()==uploaded_original_text
            assert request('/profiles/'+profile_id)==profile_before_comparison
            actions.append({'action':'Uploaded original comparison preserves revision/hash and profile facts, displays changed summary and omitted content, and downloads exact source bytes','result':'PASS'})

            select_data_job()
            page.get_by_label('Order approved skills for this job',exact=True).check()
            page.get_by_label('Choose resume facts',exact=True).check()
            expect(page.get_by_label('Include SQL',exact=True)).to_be_checked()
            page.get_by_label('Include SQL',exact=True).uncheck()
            prepare()
            selected_documents=request('/profiles/'+profile_id+'/applications/'+identifier+'/documents')['documents']
            selected_resume=next(row for row in selected_documents if row['version']==4 and row['kind']=='resume')
            selected_cover=next(row for row in selected_documents if row['version']==4 and row['kind']=='cover-letter')
            assert selected_resume['generation_model']=='deterministic-fact-selection-and-bullet-emphasis-v1'
            assert selected_resume['fact_ids']==[python_fact_id] and set(selected_cover['fact_ids'])=={python_fact_id,sql_fact_id}
            assert selected_resume['generation_context']['omitted_fact_ids']==[sql_fact_id]
            selected_history=page.get_by_role('region',name='resume version 4',exact=True)
            expect(selected_history.get_by_text('Tailored resume: selected facts',exact=False)).to_be_visible()
            with page.expect_download() as selected_download:
                selected_history.get_by_role('link',name='Download TXT',exact=True).click()
            selected_text=Path(selected_download.value.path()).read_text(encoding='utf-8')
            assert 'Python' in selected_text and 'SQL' not in selected_text
            assert 'SQL' in download_document(selected_cover['id'],'txt').decode('utf-8')
            comparison=page.get_by_role('region',name='Compare base and tailored resumes',exact=True)
            comparison.get_by_label('Base resume',exact=True).select_option(label='v2')
            comparison.get_by_role('button',name='Compare selected resumes',exact=True).click()
            expect(comparison.get_by_role('region',name='Content omitted',exact=True)).to_contain_text('skill: SQL')
            assert request('/profiles/'+profile_id)==profile_before_comparison
            actions.append({'action':'Explicit approved resume selection reaches the real document worker, omits SQL only from the resume, preserves profile and cover letter, and shows immutable omissions','result':'PASS'})

            navigation('Profile')
            page.get_by_label('city',exact=True).fill('Tulsa')
            save_profile()
            saved_contact_profile=request('/profiles/'+profile_id)
            select_data_job()
            page.get_by_label('Order approved skills for this job',exact=True).uncheck()
            contact_option=page.get_by_label('Include saved resume contact details',exact=True)
            expect(contact_option).not_to_be_checked();contact_option.check()
            prepare()
            contact_resume=next(row for row in request('/profiles/'+profile_id+'/applications/'+identifier+'/documents')['documents'] if row['version']==5 and row['kind']=='resume')
            assert contact_resume['template']=='approved-contact-base-v1'
            assert 'Location: Tulsa' in contact_resume['generation_context']['contact_header']
            contact_history=page.get_by_role('region',name='resume version 5',exact=True)
            contact_history.get_by_text('Generation details and facts',exact=True).click()
            expect(contact_history.get_by_role('region',name='Saved resume contact header')).to_contain_text('Location: Tulsa')
            with page.expect_download() as contact_download:
                contact_history.get_by_role('link',name='Download TXT',exact=True).click()
            assert 'Location: Tulsa' in Path(contact_download.value.path()).read_text(encoding='utf-8')
            assert 'Location: Tulsa' not in download_document(selected_resume['id'],'txt').decode('utf-8')
            assert request('/profiles/'+profile_id)==saved_contact_profile
            actions.append({'action':'Explicit saved contact header reaches real document worker, history and download while prior revision and profile remain unchanged','result':'PASS'})

            if onboarding:
                # Keep the uncertain-answer recovery journey specific to the onboarding fixture.
                previous_operation_id = None
                page.get_by_role("button", name="Run dry-run", exact=True).click()
                expect(page.get_by_role("status", name="Application state", exact=True)).to_have_text(
                    "REVIEW_REQUIRED", timeout=30000)
                preflight = page.get_by_role("region", name="ATS question preflight", exact=True)
                expect(preflight.get_by_role("list", name="Unresolved required questions")).to_contain_text(
                    "Approved fixture answer")
                previous_profile = request("/profiles/" + profile_id)
                review_row = preflight.get_by_role("list", name="Unresolved required questions").locator(
                    "li").filter(has_text="Approved fixture answer")
                review_row.get_by_role("button", name="Review this question", exact=True).click()
                answer_group = page.get_by_role("region", name="Application answers", exact=True).get_by_role(
                    "group", name="Question 1")
                answer_group.get_by_role("textbox", name="Answer", exact=True).fill(
                    "Controlled fictional test application, reviewed for this job.")
                answer_group.get_by_label("I reviewed this exact answer for this job", exact=True).check()
                answer_group.get_by_role("button", name="Save to Answer Bank", exact=True).click()
                answer_confirmation = answer_bank_confirmation_locator(answer_group, 1)
                answer_confirmation.get_by_label(
                    "I confirm this answer is accurate and may be reused", exact=True).check()
                with page.expect_response(lambda response:
                        response.request.method == "PUT"
                        and urlsplit(response.url).path == "/profiles") as reviewed_answer_save:
                    answer_confirmation.get_by_role(
                        "button", name="Confirm and save to Answer Bank", exact=True).click()
                assert reviewed_answer_save.value.status == 200, reviewed_answer_save.value.text()
                updated_profile = request("/profiles/" + profile_id)
                reviewed_answer = next(row for row in updated_profile["answers"]
                                       if row["question"] == "Approved fixture answer")
                assert reviewed_answer["text"] == "Controlled fictional test application, reviewed for this job."
                assert reviewed_answer["verified"] and reviewed_answer["user_approved"]
                old_pins = request("/profiles/" + profile_id + "/applications/" + identifier)
                assert old_pins["profile_version"] == previous_profile["version"]
                old_document_version = old_pins["document_version"]
                with page.expect_response(lambda candidate:
                        candidate.request.method == "POST"
                        and urlsplit(candidate.url).path.endswith("/reprepare")) as reprepare_response:
                    page.get_by_role("button", name="Resume and prepare current revision", exact=True).click()
                assert reprepare_response.value.status == 200, reprepare_response.value.text()
                expect(page.get_by_role("status").filter(
                    has_text="Run a local dry run to create a fresh ATS question preflight")).to_be_visible()
                resumed = request("/profiles/" + profile_id + "/applications/" + identifier)
                assert_reprepared_profile_pins(
                    previous_profile_version=previous_profile["version"],
                    old_application_profile_version=old_pins["profile_version"],
                    current_profile_version=updated_profile["version"],
                    resumed_application_profile_version=resumed["profile_version"],
                )
                assert resumed["document_version"] == old_document_version + 1
                assert resumed["id"] == identifier and resumed["profile_id"] == profile_id
                revised_documents = request("/profiles/" + profile_id + "/applications/" + identifier + "/documents")["documents"]
                assert any(row["version"] == old_document_version for row in revised_documents)
                assert any(row["version"] == resumed["document_version"]
                           and row["profile_version"] == updated_profile["version"]
                           for row in revised_documents)
                assert resumed.get("question_preflight") is not None
                assert resumed["question_preflight"]["is_current"] is False
                actions.append({"action": "Unknown required answer stopped dry-run; explicit question review changed approved profile facts, then explicit Resume prepared a new immutable document revision and retained old history", "result": "PASS", "boundary": "Connected React UI → public profile/reprepare APIs → actual document worker"})

            if onboarding:
                mock_controls = page.get_by_role("region", name="Recoverable mock run", exact=True)
                run_storage_key = f"mjobs-mock-run-v1:{profile_id}:{identifier}"
                previous_intent = page.evaluate(
                    "key => {const saved=localStorage.getItem(key);return saved?JSON.parse(saved):null}",
                    run_storage_key,
                )
                assert previous_intent and previous_intent["submit"] is False, (
                    "The initial review-blocked dry-run intent was not preserved for explicit resolution")
                previous_receipt = request(
                    f"/profiles/{profile_id}/applications/{identifier}/mock-runs/"
                    + previous_intent["requestId"])
                assert previous_receipt["state"] == "BLOCKED"
                assert previous_receipt["error_code"] == "REVIEW_REQUIRED"
                previous_operation_id = previous_receipt["operation_id"]
                previous_operation = request("/operations/" + previous_operation_id)
                assert_blocked_mock_operation(previous_operation, application_id=identifier)
                assert "operation_type" not in previous_operation, previous_operation
                assert resumed["question_preflight"]["is_current"] is False

                # Reprepare preserves the old terminal intent and durable receipt.
                # Require the UI's explicit user action before another run can start.
                new_run = mock_controls.get_by_role(
                    "button", name="Start a new reviewed run", exact=True)
                expect(new_run).to_be_enabled()
                new_run.click()
                assert page.evaluate("key => localStorage.getItem(key)", run_storage_key) is None
                assert request("/operations/" + previous_operation_id)["state"] == "BLOCKED"
                assert request(
                    f"/profiles/{profile_id}/applications/{identifier}/mock-runs/"
                    + previous_intent["requestId"])["operation_id"] == previous_operation_id
                expect(mock_controls.get_by_role("button", name="Run dry-run", exact=True)).to_be_enabled()
                expect(mock_controls.get_by_role("button", name="Submit to local mock", exact=True)).to_be_disabled()
                actions.append({"action": "Explicit new reviewed run clears only the prior terminal browser intent; blocked operation and stale preflight remain readable", "result": "PASS", "boundary": "Connected React UI → public run receipt/operation APIs"})

            page.get_by_role("button", name="Run dry-run", exact=True).click()
            expect(page.get_by_role("status", name="Application state", exact=True)).to_have_text(
                "DRY_RUN_COMPLETE", timeout=30000)
            dry = request("/profiles/" + profile_id + "/applications/" + identifier)
            assert "STOP_BEFORE_SUBMIT" in dry["trace"] and "SUBMIT" not in dry["trace"]
            assert dry["question_preflight"]["result"] == "PASS"
            assert dry["question_preflight"]["is_current"] is True
            assert_previous_review_operation(
                request, onboarding=onboarding,
                operation_id=previous_operation_id if onboarding else None,
                application_id=identifier if onboarding else None)
            page.get_by_role("button", name="Start a new reviewed run", exact=True).click()
            page.get_by_label("I approve one fictional mock submission").check()
            mock_run_path = f"/profiles/{profile_id}/applications/{identifier}/mock-run"
            with page.expect_response(lambda response: response.request.method == "POST"
                                      and response.url == base_url + mock_run_path) as admission:
                journal.begin("Submit to local mock")
                page.get_by_role("button", name="Submit to local mock", exact=True).click()
            assert admission.value.status == 202
            admitted = admission.value.json()
            admitted_intent = page.evaluate("([profile, application]) => JSON.parse(localStorage.getItem('mjobs-mock-run-v1:' + encodeURIComponent(profile) + ':' + encodeURIComponent(application)))", [profile_id, identifier])
            assert admitted_intent["submit"] is True
            failure_context = {
                "application_id": identifier,
                "request_id": admitted_intent.get("requestId"),
                "expected_profile_version": dry.get("profile_version"),
                "expected_document_version": dry.get("document_version"),
                "application_state": dry.get("state"),
                "preflight_result": (dry.get("question_preflight") or {}).get("result"),
                "preflight_is_current": (dry.get("question_preflight") or {}).get("is_current"),
                "submit": admitted_intent.get("submit"),
            }
            wait_for_mock_run(
                request,
                mock_run_path + "s/" + admitted_intent["requestId"],
                admitted["operation_id"],
                failure_context=failure_context,
            )
            expect(page.get_by_role("status", name="Application state", exact=True)).to_have_text("CONFIRMED")
            confirmed = request("/profiles/" + profile_id + "/applications/" + identifier)
            assert confirmed["target"] == "MOCK" and confirmed["confirmation"]["submission_count"] == 1
            assert request("/health")["external_submission_enabled"] is False
            controls = page.get_by_role('region', name='Recoverable mock run', exact=True)
            expect(controls.get_by_role('status').filter(has_text='Operation')).to_contain_text('COMPLETED')
            intent = page.evaluate("([profile, application]) => JSON.parse(localStorage.getItem('mjobs-mock-run-v1:'+encodeURIComponent(profile)+':'+encodeURIComponent(application)))", [profile_id, identifier])
            assert intent['submit'] is True
            repeated_response = context.request.post(base_url + "/profiles/" + profile_id + "/applications/" + identifier + "/mock-run", data={"submit": True, "request_id": intent['requestId']})
            assert repeated_response.status == 202
            repeated = repeated_response.json()
            assert repeated["state"] == "COMPLETED"
            assert repeated['replayed'] is True
            assert request("/profiles/" + profile_id + "/applications/" + identifier)["confirmation"] == confirmed["confirmation"]
            capture_redacted_screenshot(page, artifacts / "application-confirmed.png")
            actions.append({"action": "Browser dry-run stops before submission; explicit local mock receipt count remains one on repeat", "result": "PASS"})

            page.reload(wait_until="domcontentloaded")
            expect(page.get_by_label("Active profile")).to_have_value(profile_id)
            navigation("Applications")
            page.get_by_role("button").filter(has=page.get_by_role("heading", name="Data Engineer", exact=True)).click()
            expect(page.get_by_role("status", name="Application state", exact=True)).to_have_text("CONFIRMED")
            navigation('Profile')
            base_resumes_before_library=request('/profiles/'+profile_id+'/resumes')['resumes']
            library=page.get_by_role('region', name='Document library', exact=True)
            library.get_by_label('Document type', exact=True).select_option('CERTIFICATION')
            library.get_by_label('Document label', exact=True).fill('Synthetic certificate')
            library.get_by_label('Choose a document', exact=False).set_input_files(
                {'name':'synthetic-certificate.txt','mimeType':'text/plain','buffer':b'Synthetic library proof, not a career credential.'})
            library.get_by_role('button', name='Store document', exact=True).click()
            collection=library.get_by_role('region',name='Synthetic certificate collection',exact=True)
            expect(collection.get_by_role('link',name='Download original version 1',exact=True)).to_be_visible()
            stored=request('/profiles/'+profile_id+'/document-library')['uploaded']
            certificate=next(row for row in stored if row['document_kind']=='CERTIFICATION')
            with page.expect_download() as download:
                collection.get_by_role('link',name='Download original version 1',exact=True).click()
            assert hashlib.sha256(Path(download.value.path()).read_bytes()).hexdigest()==certificate['sha256']
            assert request('/profiles/'+profile_id+'/resumes')['resumes']==base_resumes_before_library
            if onboarding:
                assert_onboarding_profile_facts(request('/profiles/'+profile_id), pending_fact_ids)
                assert_no_pending_facts_in_documents(
                    request('/profiles/'+profile_id+'/applications/'+identifier+'/documents')['documents'],
                    pending_fact_ids, {python_fact_id, sql_fact_id})
            else:
                assert {row['fact_id'] for row in request('/profiles/'+profile_id)['facts']}=={'python','sql'}
            actions.append({'action':'Typed document library stores/downloads checksum-backed certification without adding resume-import entries or career facts','result':'PASS'})
            navigation("Activity")
            expect(page.get_by_role("heading", name="Service health", exact=True)).to_be_visible()
            page.set_viewport_size({"width": 390, "height": 844})
            assert page.evaluate("document.documentElement.scrollWidth <= innerWidth"), "Mobile overflow"
            capture_redacted_screenshot(page, artifacts / "activity-mobile.png")
            assert not failures, failures
            assert not unsafe_requests, unsafe_requests
            actions.append({"action": "Browser reload restores selected profile and application history; responsive Activity has no horizontal overflow or browser exceptions", "result": "PASS"})
            def read_bytes(path):
                response = context.request.get(base_url + path)
                assert response.ok, 'Persistence download failed'
                return response.body()
            sources = request('/company-sources')['company_sources']
            assert sources, 'Persistence proof requires a seeded company source'
            pinned = context.request.put(base_url + '/company-sources/' + sources[0]['id'] + '/pin',
                                         data={'pinned': True})
            assert pinned.ok, 'Synthetic company source control did not persist'
            assert next(row for row in request('/company-sources')['company_sources']
                        if row['id'] == sources[0]['id'])['priority'] == 'USER_PINNED'
            local_persistence = persistence.collect(request, read_bytes, profile_id, identifier)
            if onboarding:
                # Capture the public same-profile jobs/source-copy baseline while
                # the Playwright API request context is still open. The connected
                # recovery workflow compares this exact inventory after restart.
                job_inventory = collect_public_job_inventory(request, profile_id)
        except BaseException as error:
            try:
                journal.fail(error)
            except Exception:
                # Evidence I/O must not replace the UI failure being diagnosed.
                pass
            try:
                capture_redacted_screenshot(page, artifacts / "failure.png")
            except Exception:
                # The primary UI exception remains the failure reported to the caller.
                pass
            raise
        finally:
            try:
                journal.persist()
            except Exception:
                if sys.exc_info()[0] is None:
                    raise
            context.close()
            browser.close()
    evidence = {"result": "PARTIAL" if onboarding else "PASS", "timestamp": datetime.now(timezone.utc).isoformat(),
        "command": "python -m cli.ui_e2e --base-url " + base_url + (" --onboarding" if onboarding else ""), "target": "Connected React UI / PostgreSQL / three RQ queues / actual Chromium",
        "profile_id": profile_id, "application_id": identifier, "operation_id": repeated["operation_id"], "request_id": intent['requestId'], "confirmation": confirmed["confirmation"], "external_submission_enabled": False,
        "mode": "ONBOARDING_FIXTURE" if onboarding else "FICTIONAL_DEMO",
        "local_persistence": local_persistence,
        **({"job_inventory": job_inventory} if onboarding else {}),
        "discovery_operation_id": discovery_operation_id if onboarding else None,
        "discovery_request_id": discovery_request_id if onboarding else None,
        "discovery_search_proofs": discovery_search_proofs if onboarding else [],
        "limitations": ([
            "This onboarding journey uses synthetic fixture jobs and does not prove real-provider browser discovery (criterion 11).",
            "Source identity/deduplication and live public-source proof are outside this local fixture run.",
        ] if onboarding else []),
        "actions": actions, "artifacts": [str(path) for path in artifacts.glob("*") if path.name != "failure.png"]}
    journal.finish(evidence)
    print(json.dumps(evidence, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--base-url", default="http://127.0.0.1:18700")
    parser.add_argument("--evidence", required=True)
    parser.add_argument("--onboarding", action="store_true",
                        help="Run the connected UI acceptance from a new synthetic profile in ready FIXTURE mode")
    options = parser.parse_args()
    run(options.base_url, options.evidence, onboarding=options.onboarding)
